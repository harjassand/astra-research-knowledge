#!/usr/bin/env python3
"""Descriptive reanalysis of Sereda et al. (2026), Fig. 4 source data.

The spreadsheet reports field-level values pooled over three independent
experiments for Fig. 4h. This script reproduces the descriptive means and
rescue fraction only; it does not treat fields as independent animals or
recompute inferential statistics.
"""

from __future__ import annotations

import hashlib
import json
import sys
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook


def vals(sheet, column: int, rows: range) -> list[float]:
    return [float(sheet.cell(row, column).value) for row in rows
            if sheet.cell(row, column).value is not None]


def main() -> None:
    source = Path(sys.argv[1] if len(sys.argv) > 1 else
                  "work/agents/c9_aging_control/sources/sereda2026_source_fig4.xlsx")
    raw = source.read_bytes()
    wb = load_workbook(source, data_only=True, read_only=True)
    sheet = wb["Fig. 4"]

    # Fig. 4e: three independently repeated observations per cell. Columns
    # C:E are untreated control fibroblasts and F:H are palbociclib-senescent
    # control fibroblasts; rows 126/127 are control/CMA-deficient macrophages.
    f4e = {
        "WT macrophage, untreated fibroblast": vals(sheet, 3, range(126, 127))
        + vals(sheet, 4, range(126, 127)) + vals(sheet, 5, range(126, 127)),
        "WT macrophage, senescent fibroblast": vals(sheet, 6, range(126, 127))
        + vals(sheet, 7, range(126, 127)) + vals(sheet, 8, range(126, 127)),
        "CMA-KO macrophage, untreated fibroblast": vals(sheet, 3, range(127, 128))
        + vals(sheet, 4, range(127, 128)) + vals(sheet, 5, range(127, 128)),
        "CMA-KO macrophage, senescent fibroblast": vals(sheet, 6, range(127, 128))
        + vals(sheet, 7, range(127, 128)) + vals(sheet, 8, range(127, 128)),
    }
    means4e = {k: mean(v) for k, v in f4e.items()}

    # Fig. 4h: bead-uptake fluorescence from 10/11 fields per group across
    # three independent experiments. Columns B/C are control/CMA-KO senescent
    # fibroblast secretome without CMA activator; D/E are the matched CA arms.
    f4h = {
        "control secretome, no CA": vals(sheet, 2, range(173, 184)),
        "CMA-KO secretome, no CA": vals(sheet, 3, range(173, 184)),
        "control secretome, CA": vals(sheet, 4, range(173, 184)),
        "CMA-KO secretome, CA": vals(sheet, 5, range(173, 184)),
    }
    means4h = {k: mean(v) for k, v in f4h.items()}
    baseline = means4h["control secretome, no CA"]
    inhibited = means4h["CMA-KO secretome, no CA"]
    rescued = means4h["CMA-KO secretome, CA"]
    gap = baseline - inhibited

    out = {
        "source": {
            "citation": "Sereda et al., Nature Aging (2026), DOI 10.1038/s43587-026-01240-w, Fig. 4 source data",
            "source_data_sha256": hashlib.sha256(raw).hexdigest(),
            "file": str(source),
        },
        "fig4e_apoptotic_corpse_efferocytosis": {
            "units": "source-reported percent efferocytosis",
            "raw_values": f4e,
            "means": means4e,
            "WT_to_KO_ratio_untreated_fibroblasts": means4e["WT macrophage, untreated fibroblast"] / means4e["CMA-KO macrophage, untreated fibroblast"],
            "WT_to_KO_ratio_senescent_fibroblasts": means4e["WT macrophage, senescent fibroblast"] / means4e["CMA-KO macrophage, senescent fibroblast"],
            "relative_fibroblast_state_change_with_WT_macrophage": (means4e["WT macrophage, senescent fibroblast"] - means4e["WT macrophage, untreated fibroblast"]) / means4e["WT macrophage, untreated fibroblast"],
            "relative_fibroblast_state_change_with_KO_macrophage": (means4e["CMA-KO macrophage, senescent fibroblast"] - means4e["CMA-KO macrophage, untreated fibroblast"]) / means4e["CMA-KO macrophage, untreated fibroblast"],
            "source_anova_percent_total_variation": {
                "macrophage_genotype": 98.22,
                "fibroblast_state": 0.8562,
                "interaction": 0.7214,
            },
            "replication_unit": "n=3 independent experiments; descriptive group means only",
        },
        "fig4h_bead_uptake_fluorescence": {
            "units": "source-reported FITC-IgG fluorescence, not direct senescent-cell clearance",
            "raw_field_values": f4h,
            "means": means4h,
            "rescue_fraction_of_no_CA_secretome_gap": (rescued - inhibited) / gap,
            "remaining_fraction_of_no_CA_secretome_gap_after_CA": (baseline - rescued) / gap,
            "field_counts": {k: len(v) for k, v in f4h.items()},
            "replication_unit": "10/11 fields per group nested in three independent experiments; no field-level CI is reported here",
        },
        "limits": [
            "Fig. 4e measures efferocytosis of staurosporine-killed targets, not clearance of living senescent cells.",
            "Fig. 4h measures macrophage uptake of IgG-coated beads after secretome exposure, not target-cell fate.",
            "Fields are nested within experiments; field counts are not treated as independent biological replicates.",
            "No in-vivo per-cell CMA activity, senescent-cell birth flux, macrophage survival flux, or clearance hazard is inferred.",
        ],
    }
    print(json.dumps(out, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
