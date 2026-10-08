#!/usr/bin/env python3
"""Audit the public CD36 rate-calculation input interface.

Requires openpyxl (available in the Codex bundled Python runtime). This reads
only the two small source workbooks selectively extracted by HTTP byte ranges;
it does not download the 2.5-GB archive.
"""

from __future__ import annotations

import json
import math
import statistics
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"


def infer_rates(f_on: float, odds_ratio: float, birth_rate: float) -> dict[str, float]:
    c = f_on * (1.0 - f_on)
    m = 2.0 * birth_rate / (math.sqrt(1.0 + 4.0 * c * (odds_ratio - 1.0)) - 1.0)
    return {"f_on": f_on, "k_on": f_on * m, "k_off": (1.0 - f_on) * m, "k_sum": m, "memory_time_h": 1.0 / m}


def main() -> None:
    frac_wb = load_workbook(DATA / "IL1 and CD36 fractions repository.xlsx", read_only=True, data_only=True)
    ws = frac_wb["Cd36 fractions"]
    clone_rows = []
    for n_log2, fraction in ws.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True):
        if isinstance(n_log2, (int, float)) and isinstance(fraction, (int, float)):
            clone_rows.append((float(n_log2), min(max(float(fraction), 1e-4), 1.0 - 1e-4)))
    f_clone_unweighted = statistics.mean(f for _, f in clone_rows)
    f_clone_cell_weighted = sum((2.0**n) * f for n, f in clone_rows) / sum(2.0**n for n, _ in clone_rows)

    pair_wb = load_workbook(DATA / "odds_ratio_CD36_tnfa_il1b.xlsx", read_only=True, data_only=True)
    pair_ws = pair_wb["Sheet1"]
    plate_rows = []
    for row in pair_ws.iter_rows(min_row=2, values_only=True):
        for start in (8, 16):
            if row[start] is None or row[start + 1] != "per_cd36_intensity_80_mean" or row[start + 2] != "Mock":
                continue
            plate, on_on, off_off, mixed = int(row[start]), int(row[start + 3]), int(row[start + 4]), int(row[start + 5])
            pairs = on_on + off_off + mixed
            f_pair = (2 * on_on + mixed) / (2 * pairs)
            odds_ratio = 4 * on_on * off_off / (mixed * mixed)
            plate_rows.append({"plate": plate, "on_on": on_on, "off_off": off_off, "mixed": mixed,
                               "pairs": pairs, "paired_cell_marginal": f_pair, "odds_ratio": odds_ratio})

    a = sum(r["on_on"] for r in plate_rows)
    b = sum(r["off_off"] for r in plate_rows)
    c = sum(r["mixed"] for r in plate_rows)
    pairs = a + b + c
    f_pair_aggregate = (2 * a + c) / (2 * pairs)
    or_aggregate = 4 * a * b / (c * c)
    birth_rate = math.log(2.0) / 14.4
    result = {
        "source_input": {
            "fraction_workbook_sheet": "Cd36 fractions",
            "numeric_clone_rows": len(clone_rows),
            "R_script_fraction_source": "mean(clone-level FractionON), clamped to [1e-4, 1-1e-4]",
            "pair_workbook_dataset": "per_cd36_intensity_80_mean / Mock (the dataset5 selection in the Rmd)",
            "pair_counts": {"ON_ON": a, "OFF_OFF": b, "mixed": c, "pairs": pairs},
            "birth_rate_per_hour": birth_rate,
        },
        "plate_rows": plate_rows,
        "fraction_comparison": {
            "R_script_unweighted_clone_fraction": f_clone_unweighted,
            "cell_size_weighted_clone_fraction_diagnostic": f_clone_cell_weighted,
            "paired_cell_marginal_same_pair_table": f_pair_aggregate,
            "difference_script_minus_paired_marginal": f_clone_unweighted - f_pair_aggregate,
        },
        "same_OR_rate_inversion": {
            "OR": or_aggregate,
            "using_R_script_clone_mean": infer_rates(f_clone_unweighted, or_aggregate, birth_rate),
            "using_pair_table_marginal": infer_rates(f_pair_aggregate, or_aggregate, birth_rate),
            "warning": "This isolates the f-input mismatch algebraically. It is not a replay of the paper's full bootstrap or proof that the authors used this exact script version for every reported parameter."
        },
    }
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
