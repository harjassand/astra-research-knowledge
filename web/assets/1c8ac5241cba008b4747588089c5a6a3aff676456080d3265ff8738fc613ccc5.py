#!/usr/bin/env python3
"""Read-only structural inventory for Dang et al. source-data workbook.

This script never writes to the source workbook. It emits a compact JSON map of
sheet names, dimensions, populated rows, formula counts, formats, and cell
samples so the experimental record grain can be reconstructed before analysis.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def inspect(path: Path) -> dict[str, Any]:
    raw = path.read_bytes()
    wb = load_workbook(path, data_only=False, read_only=False)
    result: dict[str, Any] = {
        "source_path": str(path.resolve()),
        "sha256": hashlib.sha256(raw).hexdigest(),
        "file_bytes": len(raw),
        "workbook_properties": {
            "creator": wb.properties.creator,
            "title": wb.properties.title,
            "subject": wb.properties.subject,
            "created": json_value(wb.properties.created),
            "modified": json_value(wb.properties.modified),
            "sheet_names": wb.sheetnames,
        },
        "sheets": [],
    }
    for ws in wb.worksheets:
        populated_rows = 0
        sample_rows = []
        nonempty_cells = 0
        formula_count = 0
        formats: dict[str, int] = {}
        nonempty_columns: set[int] = set()
        for row in ws.iter_rows():
            cells = []
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                nonempty_cells += 1
                nonempty_columns.add(cell.column)
                formats[cell.number_format] = formats.get(cell.number_format, 0) + 1
                entry = {
                    "cell": cell.coordinate,
                    "value": json_value(value),
                    "type": cell.data_type,
                }
                if cell.data_type == "f":
                    formula_count += 1
                cells.append(entry)
            if cells:
                populated_rows += 1
                if row[0].row <= 12:
                    # Keep the map compact for wide source matrices.
                    sample_rows.append({
                        "row": row[0].row,
                        "cells": cells[:12] + (cells[-4:] if len(cells) > 16 else cells[12:]),
                    })
        max_pop_row = max((r["row"] for r in sample_rows), default=0)
        # Recompute the last populated row from the worksheet extent and the
        # nonempty-row counter is intentionally not a maximum-row proxy.
        for row_num in range(ws.max_row, 0, -1):
            if any(ws.cell(row_num, col).value is not None for col in range(1, ws.max_column + 1)):
                max_pop_row = row_num
                break
        max_pop_col = max(nonempty_columns, default=0)
        result["sheets"].append({
            "name": ws.title,
            "state": ws.sheet_state,
            "declared_dimension": ws.calculate_dimension(),
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "populated_dimension": f"A1:{ws.cell(max_pop_row, max_pop_col).coordinate if max_pop_row and max_pop_col else 'A1'}",
            "nonempty_cells": nonempty_cells,
            "populated_row_count": populated_rows,
            "sample_rows": sample_rows,
            "formula_count": formula_count,
            "number_format_counts": formats,
            "merged_ranges": [str(x) for x in ws.merged_cells.ranges],
            "tables": [
                {"name": t.displayName, "ref": t.ref}
                for t in ws.tables.values()
            ],
            "freeze_panes": str(ws.freeze_panes) if ws.freeze_panes else None,
            "auto_filter": ws.auto_filter.ref,
        })
    wb.close()
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    result = inspect(args.workbook)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps({
        "sha256": result["sha256"],
        "file_bytes": result["file_bytes"],
        "sheets": [
            {key: s[key] for key in ("name", "state", "declared_dimension", "max_row", "max_column", "nonempty_cells", "populated_row_count", "formula_count", "number_format_counts")}
            for s in result["sheets"]
        ],
        "map_path": str(args.out.resolve()),
    }, indent=2))


if __name__ == "__main__":
    main()
