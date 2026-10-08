#!/usr/bin/env python3
"""Read the published Nature Fig. 3 source workbook and run model diagnostics.

Uses the bundled Python/openpyxl runtime documented in source_data_metadata.md.
The workbook is read-only; outputs are deterministic JSON from supplied
replicate-level percentages, not event-level flow-cytometry data.
"""

from __future__ import annotations

import json
import math
import statistics
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parents[1]
WORKBOOK = HERE / "data" / "source_data_fig3.xlsx"
OUTPUT = HERE / "source_data_analysis.json"


def mean_percent(sheet, rows: range, col: int) -> float:
    return statistics.mean(float(sheet.cell(r, col).value) for r in rows) / 100.0


def source_data_diagnostics() -> dict:
    wb = load_workbook(WORKBOOK, data_only=True, read_only=True)

    # Fig. 3a-c: single-device marginal cultures, followed by the joint 2-device
    # four-state culture. States: GD, RD, RG, DD (G=Venus, R=mScarlet).
    s2 = wb["Figure 3 a-c"]
    p_g = mean_percent(s2, range(9, 17), 5)
    p_r = mean_percent(s2, range(9, 17), 12)
    labels2 = ["GD", "RD", "RG", "DD"]
    cols2 = [5, 6, 7, 8]
    observed2 = [mean_percent(s2, range(32, 40), c) for c in cols2]
    independent2 = [
        p_g * (1 - p_r),
        (1 - p_g) * p_r,
        p_g * p_r,
        (1 - p_g) * (1 - p_r),
    ]
    dual_p_g = observed2[0] + observed2[2]
    dual_p_r = observed2[1] + observed2[2]
    dual_cov = observed2[2] - dual_p_g * dual_p_r

    # Fig. 3e-g: three single-device marginals and the joint 3-device, 8-state
    # culture. Bit order is (G,R,B); state labels are those in the workbook.
    s3 = wb["Figure 3 e-g"]
    marginal3 = [
        mean_percent(s3, range(12, 20), 6),
        mean_percent(s3, range(12, 20), 14),
        mean_percent(s3, range(12, 20), 21),
    ]
    labels3 = ["DDD", "GDD", "RDD", "BDD", "RGD", "GBD", "RBD", "RGB"]
    state_bits = {
        "DDD": (0, 0, 0),
        "GDD": (1, 0, 0),
        "RDD": (0, 1, 0),
        "BDD": (0, 0, 1),
        "RGD": (1, 1, 0),
        "GBD": (1, 0, 1),
        "RBD": (0, 1, 1),
        "RGB": (1, 1, 1),
    }
    observed3 = [mean_percent(s3, range(44, 52), c) for c in range(5, 13)]
    independent3 = [
        math.prod(p if bit else 1 - p for p, bit in zip(marginal3, state_bits[label]))
        for label in labels3
    ]
    joint_marginal3 = [
        sum(q for label, q in zip(labels3, observed3) if state_bits[label][i])
        for i in range(3)
    ]
    pairwise3 = {}
    for i, j, name in [(0, 1, "G-R"), (0, 2, "G-B"), (1, 2, "R-B")]:
        q_ij = sum(
            q for label, q in zip(labels3, observed3)
            if state_bits[label][i] and state_bits[label][j]
        )
        pairwise3[name] = {
            "joint_probability": q_ij,
            "product_of_joint_marginals": joint_marginal3[i] * joint_marginal3[j],
            "covariance": q_ij - joint_marginal3[i] * joint_marginal3[j],
        }

    return {
        "provenance": {
            "source_doi": "10.1038/s41586-026-10259-3",
            "source_data_item": "Source Data Fig. 3 (MOESM11_ESM.xlsx)",
            "replicates": 8,
            "measurement": "published replicate-level flow-cytometry percentages",
            "event_level_fcs_available_in_source_data_item": False,
        },
        "two_parallel_modules": {
            "single_module_marginals": {"G": p_g, "R": p_r},
            "independence_prediction_state_order_GD_RD_RG_DD": independent2,
            "observed_joint_state_order_GD_RD_RG_DD": observed2,
            "state_order": labels2,
            "total_variation_from_single_module_product": 0.5 * sum(
                abs(a - b) for a, b in zip(independent2, observed2)
            ),
            "joint_culture_marginals": {"G": dual_p_g, "R": dual_p_r},
            "joint_culture_covariance_G_R": dual_cov,
        },
        "three_parallel_modules": {
            "single_module_marginals_G_R_B": marginal3,
            "independence_prediction_state_order": independent3,
            "observed_joint_state_order": observed3,
            "state_order": labels3,
            "total_variation_from_single_module_product": 0.5 * sum(
                abs(a - b) for a, b in zip(independent3, observed3)
            ),
            "joint_culture_marginals_G_R_B": joint_marginal3,
            "within_joint_culture_pairwise_covariances": pairwise3,
        },
    }


def shared_exposure_all_on_example(m: int = 10, active_fraction: float = 0.1,
                                   active_probability: float = 0.99) -> dict:
    """Exact two-dose example showing the independent product error."""
    marginal = active_fraction * active_probability
    true_all_on = active_fraction * active_probability**m
    product_all_on = marginal**m
    return {
        "modules": m,
        "exposure_law": {"A=0_probability": 1 - active_fraction,
                         "A=high_probability": active_fraction},
        "per_module_activation_given_A0": 0.0,
        "per_module_activation_given_Ahigh": active_probability,
        "single_module_marginal": marginal,
        "shared_exposure_all_on_probability": true_all_on,
        "product_of_marginals_all_on_prediction": product_all_on,
        "underprediction_factor": true_all_on / product_all_on,
    }


def one_hot_target_parallel_obstruction(m: int) -> dict:
    """Association witness separating a one-hot target from common-dose products."""
    p = 1.0 / m
    # A one-hot target has P(X_i=X_j=1)=0, whereas common-dose mixtures of
    # conditionally independent, monotone module responses obey q_ij >= p_i p_j.
    return {
        "fates": m,
        "target_marginal_per_module": p,
        "target_pairwise_joint_probability": 0.0,
        "minimum_common_exposure_pairwise_joint_probability": p * p,
        "pairwise_covariance_gap": -(p * p),
        "separation_witness": "P(X_i=1,X_j=1) - P(X_i=1)P(X_j=1) >= 0",
    }


def huffman_tree(probabilities: dict[str, float]) -> tuple[dict, float]:
    """Compile a categorical target to a minimum-mean-depth binary tree."""
    if not probabilities or any(p <= 0 for p in probabilities.values()):
        raise ValueError("All listed target fates must have positive probability")
    total = sum(probabilities.values())
    if not math.isclose(total, 1.0, rel_tol=0, abs_tol=1e-12):
        raise ValueError("Target fate probabilities must sum to one")

    # Deterministic Huffman tie-breaking by monotonically assigned node id.
    heap: list[tuple[float, int, dict]] = []
    next_id = 0
    for label, p in sorted(probabilities.items()):
        heap.append((p, next_id, {"fate": label, "mass": p}))
        next_id += 1
    import heapq
    heapq.heapify(heap)
    if len(heap) == 1:
        return heap[0][2], 0.0
    while len(heap) > 1:
        p0, _, n0 = heapq.heappop(heap)
        p1, _, n1 = heapq.heappop(heap)
        node = {"mass": p0 + p1, "left": n0, "right": n1,
                "left_probability": p0 / (p0 + p1),
                "right_probability": p1 / (p0 + p1)}
        heapq.heappush(heap, (p0 + p1, next_id, node))
        next_id += 1
    root = heap[0][2]

    def assign(node: dict, depth: int, out: dict[str, float]) -> None:
        if "fate" in node:
            out[node["fate"]] = probabilities[node["fate"]]
            node["depth"] = depth
            return
        node["depth"] = depth
        assign(node["left"], depth + 1, out)
        assign(node["right"], depth + 1, out)

    assign(root, 0, {})
    mean_depth = sum(probabilities[label] * _leaf_depth(root, label) for label in probabilities)
    return root, mean_depth


def _leaf_depth(node: dict, label: str, depth: int = 0) -> int:
    if "fate" in node:
        if node["fate"] != label:
            raise KeyError(label)
        return depth
    for side in ("left", "right"):
        try:
            return _leaf_depth(node[side], label, depth + 1)
        except KeyError:
            pass
    raise KeyError(label)


def _leaf_probabilities(node: dict, path_probability: float = 1.0) -> dict[str, float]:
    if "fate" in node:
        return {node["fate"]: path_probability}
    left = _leaf_probabilities(node["left"], path_probability * node["left_probability"])
    right = _leaf_probabilities(node["right"], path_probability * node["right_probability"])
    return {**left, **right}


def main() -> None:
    out = source_data_diagnostics()
    out["synthetic_counterexample"] = shared_exposure_all_on_example()
    out["one_hot_10_fate_obstruction"] = one_hot_target_parallel_obstruction(10)
    tree, mean_depth = huffman_tree({f"fate_{i+1}": 0.1 for i in range(10)})
    compiled = _leaf_probabilities(tree)
    target = {f"fate_{i+1}": 0.1 for i in range(10)}
    epsilon, delta, dose_lipschitz = 0.05, 0.05, 10.0
    full_law_sensor_n = math.ceil(
        dose_lipschitz**2 * math.log(2 / delta) / (2 * epsilon**2)
    )
    out["sequential_compiler_example"] = {
        "target": target,
        "tree": tree,
        "binary_modules": 9,
        "mean_branch_decisions": mean_depth,
        "reconstructed_leaf_probabilities": compiled,
        "max_absolute_compilation_error": max(abs(compiled[k] - target[k]) for k in target),
    }
    out["polynomial_assay_bound_example"] = {
        "m": 10,
        "Amax_times_sum_response_lipschitz": dose_lipschitz,
        "whole_law_total_variation_tolerance": epsilon,
        "confidence": 1 - delta,
        "independent_single_cell_exposure_measurements_required": full_law_sensor_n,
        "scope": "dose-sensor sample only; add response-curve, sensor-calibration, and biological-replicate costs",
    }
    OUTPUT.write_text(json.dumps(out, indent=2) + "\n")
    print(json.dumps(out, indent=2))


if __name__ == "__main__":
    main()
