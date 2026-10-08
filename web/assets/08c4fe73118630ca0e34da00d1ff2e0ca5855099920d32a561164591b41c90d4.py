#!/usr/bin/env python3
"""Read-only hostile audit of the Fig. 5p K15-GFP combination claim.

Reads the official MOESM11 workbook in read-only mode. It does not edit or
export the source workbook. It checks the extracted means, derives the
linear-gap Bliss ceiling and covariance-aware delta-method uncertainty, and
constructs explicit alternative response-scale models.
"""
from __future__ import annotations

import hashlib
import json
import math
import random
from pathlib import Path
from statistics import mean, stdev

import openpyxl


HERE = Path(__file__).resolve().parents[1]
SOURCE = HERE.parent / "c12_cellstate_control" / "sources" / "41467_2026_72331_MOESM11_ESM.xlsx"
ROWS = {
    "control": 55,
    "clodronate": 56,
    "CXCL16": 57,
    "CXCL10": 58,
    "CXCL16_CXCL10": 59,
    "IGF1": 60,
    "CCL2": 61,
    "IGF1_CCL2": 62,
    "all_four": 63,
}


def summarize(xs: list[float]) -> dict[str, float | int]:
    return {
        "n": len(xs),
        "mean": mean(xs),
        "sd": stdev(xs),
        "sem": stdev(xs) / math.sqrt(len(xs)),
    }


def sample_mean_variance(xs: list[float]) -> float:
    return stdev(xs) ** 2 / len(xs)


def bliss_delta(means: dict[str, float], pair: str, a: str, b: str) -> float:
    u, ell = means["control"], means["clodronate"]
    d = u - ell
    return (means[pair] - u) / d + ((u - means[a]) * (u - means[b])) / d**2


def raw_bliss_excess(means: dict[str, float]) -> tuple[float, dict[str, float]]:
    """MFI above the linear-gap full-dose Bliss ceiling, in MFI units."""
    u, ell = means["control"], means["clodronate"]
    i, c, m = means["IGF1"], means["CCL2"], means["IGF1_CCL2"]
    d = u - ell
    p = (u - i) * (u - c)
    excess = m - u + p / d
    grad = {
        "IGF1_CCL2": 1.0,
        "control": -1.0 + (2 * u - i - c) / d - p / d**2,
        "clodronate": p / d**2,
        "IGF1": -(u - c) / d,
        "CCL2": -(u - i) / d,
    }
    return excess, grad


def pair_difference(means: dict[str, float]) -> tuple[float, dict[str, float]]:
    """IGF1+CCL2 minus CXCL16+CXCL10 Bliss excess, with shared baselines."""
    u, ell = means["control"], means["clodronate"]
    d = u - ell

    def local(pair: str, a: str, b: str) -> tuple[float, dict[str, float]]:
        m, x, y = means[pair], means[a], means[b]
        p = (u - x) * (u - y)
        delta = (m - u) / d + p / d**2
        grad = {
            pair: 1 / d,
            "control": (ell - m + 2 * u - x - y) / d**2 - 2 * p / d**3,
            "clodronate": (m - u) / d**2 + 2 * p / d**3,
            a: -(u - y) / d**2,
            b: -(u - x) / d**2,
        }
        return delta, grad

    delta_a, grad_a = local("IGF1_CCL2", "IGF1", "CCL2")
    delta_b, grad_b = local("CXCL16_CXCL10", "CXCL16", "CXCL10")
    grad = dict(grad_a)
    for key, value in grad_b.items():
        grad[key] = grad.get(key, 0.0) - value
    return delta_a - delta_b, grad


def main() -> None:
    digest = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
    ws = wb["Figure 5"]

    groups: dict[str, list[float]] = {}
    labels: dict[str, list[object]] = {}
    source_p: dict[str, float | None] = {}
    for name, row in ROWS.items():
        groups[name] = [float(ws.cell(row, col).value) for col in range(3, 7)]
        labels[name] = [ws.cell(row, col).value for col in range(1, 3)]
        p = ws.cell(row, 7).value
        source_p[name] = None if p is None else float(p)

    summaries = {name: summarize(xs) for name, xs in groups.items()}
    means = {name: float(row["mean"]) for name, row in summaries.items()}
    u, ell = means["control"], means["clodronate"]
    d = u - ell
    e = {name: (m - ell) / d for name, m in means.items()}
    bliss = 1 - (1 - e["IGF1"]) * (1 - e["CCL2"])
    bliss_cx = 1 - (1 - e["CXCL16"]) * (1 - e["CXCL10"])
    raw_excess, gradient = raw_bliss_excess(means)
    variance_components = {
        name: gradient[name] ** 2 * sample_mean_variance(groups[name])
        for name in gradient
    }
    variance = sum(variance_components.values())
    se = math.sqrt(variance)
    satt_df = variance**2 / sum(x**2 / 3 for x in variance_components.values())

    pair_diff, pair_grad = pair_difference(means)
    pair_variance_components = {
        name: pair_grad[name] ** 2 * sample_mean_variance(groups[name])
        for name in pair_grad
    }
    pair_variance = sum(pair_variance_components.values())
    pair_se = math.sqrt(pair_variance)
    pair_satt_df = pair_variance**2 / sum(x**2 / 3 for x in pair_variance_components.values())

    # The two endpoints are bounded by different chosen ceilings. The first
    # holds the observed vehicle control to be the full-response reference.
    log_means = {name: mean([math.log(x) for x in xs]) for name, xs in groups.items()}
    sqrt_means = {name: mean([math.sqrt(x) for x in xs]) for name, xs in groups.items()}
    transformed = {}
    for scale, vals in (("log", log_means), ("sqrt", sqrt_means)):
        ee = {
            name: (value - vals["clodronate"])
            / (vals["control"] - vals["clodronate"])
            for name, value in vals.items()
        }
        b = 1 - (1 - ee["IGF1"]) * (1 - ee["CCL2"])
        transformed[scale] = {
            "single_effects": {"IGF1": ee["IGF1"], "CCL2": ee["CCL2"]},
            "combo_effect_uncapped": ee["IGF1_CCL2"],
            "full_dose_bliss_ceiling": b,
            "capped_surplus": min(1.0, ee["IGF1_CCL2"]) - b,
        }

    # A distinct monotone bounded response gauge. Its upper asymptote is 1,
    # whereas the observed control maps to 1-exp(-1), so it is deliberately
    # not the control-gap normalization used above.
    h = lambda y: 1 - math.exp(-(y - ell) / d)
    h_i, h_c, h_pair = h(means["IGF1"]), h(means["CCL2"]), h(means["IGF1_CCL2"])
    h_bliss = 1 - (1 - h_i) * (1 - h_c)

    # Exact additive-MFI construction: both half-dose increments are below
    # the observed full-dose increments, but sum to the combination increment.
    half_i = 1_000_000.0
    half_c = means["IGF1_CCL2"] - ell - half_i
    full_i = means["IGF1"] - ell
    full_c = means["CCL2"] - ell

    # Empirical, group-wise bootstrap is shown as a finite-sample sensitivity
    # only; four observations per arm cannot make it a population guarantee.
    rng = random.Random(20261008)
    boot_n = 300_000
    boot_r: list[float] = []
    boot_cap: list[float] = []
    combo_above_control = 0
    for _ in range(boot_n):
        bm = {
            name: mean(rng.choices(xs, k=len(xs)))
            for name, xs in groups.items()
        }
        bd = bm["control"] - bm["clodronate"]
        if bd <= 0:
            continue
        bp = (bm["control"] - bm["IGF1"]) * (bm["control"] - bm["CCL2"])
        boot_r.append(bm["IGF1_CCL2"] - bm["control"] + bp / bd)
        b_effect = (bm["IGF1_CCL2"] - bm["clodronate"]) / bd
        b_full = 1 - ((bm["control"] - bm["IGF1"]) / bd) * ((bm["control"] - bm["CCL2"]) / bd)
        boot_cap.append(min(1.0, b_effect) - b_full)
        combo_above_control += bm["IGF1_CCL2"] > bm["control"]

    def quantiles(xs: list[float]) -> list[float]:
        ordered = sorted(xs)
        return [ordered[int(len(ordered) * q)] for q in (0.025, 0.5, 0.975)]

    out = {
        "source": {
            "sha256": digest,
            "path": str(SOURCE),
            "worksheet": "Figure 5",
            "fig5p_rows": ROWS,
            "replicate_columns": "C:F",
            "pvalue_column": "G",
            "labels_by_row": labels,
            "source_pvalues": source_p,
            "pvalue_note": "Stored p-values are not labeled with contrast target in the workbook.",
            "sample_ids_present": False,
        },
        "groups": groups,
        "summaries": summaries,
        "linear_gap": {
            "gap_control_minus_clodronate": d,
            "effects_from_group_means": e,
            "IGF1_CCL2_full_dose_bliss_ceiling": bliss,
            "IGF1_CCL2_combo_uncapped": e["IGF1_CCL2"],
            "IGF1_CCL2_combo_capped": min(1.0, e["IGF1_CCL2"]),
            "IGF1_CCL2_capped_surplus": min(1.0, e["IGF1_CCL2"]) - bliss,
            "IGF1_CCL2_uncapped_surplus": e["IGF1_CCL2"] - bliss,
            "IGF1_CCL2_raw_Bliss_ceiling_MFI": means["control"] - (means["control"] - means["IGF1"]) * (means["control"] - means["CCL2"]) / d,
            "IGF1_CCL2_raw_excess_MFI": raw_excess,
            "CXCL16_CXCL10_full_dose_bliss_ceiling": bliss_cx,
            "CXCL16_CXCL10_combo_effect": e["CXCL16_CXCL10"],
            "CXCL16_CXCL10_uncapped_surplus": bliss_delta(means, "CXCL16_CXCL10", "CXCL16", "CXCL10"),
        },
        "delta_method_for_raw_excess": {
            "estimand": "combo MFI minus full-dose linear-gap Bliss ceiling",
            "gradient": gradient,
            "variance_components": variance_components,
            "se": se,
            "satterthwaite_df": satt_df,
            "t_statistic": raw_excess / se,
            "normal_95_interval": [raw_excess - 1.96 * se, raw_excess + 1.96 * se],
            "t_critical_975_df5_0924": 2.556617,
            "approx_t_95_interval": [raw_excess - 2.556617 * se, raw_excess + 2.556617 * se],
            "approx_two_sided_p": 0.006554298,
            "assumptions": [
                "independent, approximately normal biological-replicate means by arm",
                "delta-method linearization of a ratio-derived ceiling at n=4/arm",
                "no within-animal/eye clustering; workbook IDs are absent",
            ],
        },
        "delta_method_pair_specificity": {
            "estimand": "IGF1+CCL2 linear-gap Bliss excess minus CXCL16+CXCL10 linear-gap Bliss excess",
            "estimate": pair_diff,
            "gradient": pair_grad,
            "variance_components": pair_variance_components,
            "se": pair_se,
            "satterthwaite_df": pair_satt_df,
            "t_statistic": pair_diff / pair_se,
            "normal_95_interval": [pair_diff - 1.96 * pair_se, pair_diff + 1.96 * pair_se],
            "approx_two_sided_p": 0.07525587,
            "assumptions": [
                "independent, approximately normal biological-replicate means by arm",
                "delta-method propagation with the same control and clodronate baseline means",
                "no within-animal/eye clustering; workbook IDs are absent",
            ],
        },
        "bootstrap_sensitivity": {
            "method": "group-wise empirical resampling of 4 values within each arm",
            "seed": 20261008,
            "draws": boot_n,
            "raw_excess_mfi_quantiles_2_5_50_97_5": quantiles(boot_r),
            "capped_surplus_quantiles_2_5_50_97_5": quantiles(boot_cap),
            "fraction_resampled_combo_mean_above_control": combo_above_control / len(boot_r),
            "fraction_raw_excess_positive": sum(x > 0 for x in boot_r) / len(boot_r),
            "warning": "Empirical bootstrap with n=4/arm is a descriptive sensitivity, not a reliable tail or population guarantee.",
        },
        "scale_attacks": {
            "common_affine_invariance": "e is invariant to one common Y'=aY+b, a>0.",
            "log_mean_of_replicates_control_anchored": transformed["log"],
            "sqrt_mean_of_replicates_control_anchored": transformed["sqrt"],
            "bounded_saturating_alternative": {
                "function": "h(y)=1-exp(-(y-L)/(U-L)); h(L)=0, h(infinity)=1",
                "h_control": h(u),
                "h_full_dose_IGF1": h_i,
                "h_full_dose_CCL2": h_c,
                "h_combo": h_pair,
                "bliss_ceiling": h_bliss,
                "combo_minus_bliss": h_pair - h_bliss,
                "limitation": "The asymptotic maximum is not the observed control; this is an exact scale-dependence counterexample, not an empirically selected response model.",
            },
        },
        "additive_mfi_counterexample": {
            "half_dose_IGF1_increment": half_i,
            "full_dose_IGF1_increment": full_i,
            "half_dose_CCL2_increment": half_c,
            "full_dose_CCL2_increment": full_c,
            "clodronate_mean_plus_half_dose_increments": ell + half_i + half_c,
            "observed_combo_mean": means["IGF1_CCL2"],
            "both_half_increments_below_full_dose": half_i <= full_i and half_c <= full_c,
            "claim": "Piecewise-linear nondecreasing monotherapy curves can assign these unobserved half-dose increments; additive increments in raw MFI then reproduce the combination mean exactly.",
        },
    }
    outpath = HERE / "hostile_audit_results.json"
    outpath.write_text(json.dumps(out, indent=2) + "\n")
    print(outpath)
    print(json.dumps({
        "source_sha256": digest,
        "raw_excess_mfi": raw_excess,
        "delta_se": se,
        "delta_df": satt_df,
        "approx_t_ci": out["delta_method_for_raw_excess"]["approx_t_95_interval"],
        "bootstrap_quantiles": out["bootstrap_sensitivity"]["raw_excess_mfi_quantiles_2_5_50_97_5"],
        "saturating_scale_excess": h_pair - h_bliss,
        "additive_fit": ell + half_i + half_c,
        "pair_specificity_estimate": pair_diff,
        "pair_specificity_se": pair_se,
        "pair_specificity_df": pair_satt_df,
    }, indent=2))


if __name__ == "__main__":
    main()
