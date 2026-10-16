"""Own exact grid pairing and conditional balance reconstruction."""
import openpyxl,json,pathlib
from fractions import Fraction
p=pathlib.Path(__file__).parent;w=openpyxl.load_workbook(p/'final_Fig2.xlsx',data_only=True)
a,b=w['Figure 2 c'],w['Figure 2 d'];groups={}
for r in range(2,32):
 h,d,X=[a.cell(r,c).value for c in (2,3,4)];hb,db,S=[b.cell(r,c).value for c in (2,3,4)]
 assert (h,d)==(hb,db)
 duty=str(Fraction(str(h))/Fraction(str(d)))
 item={'row':r,'h_s':h,'d_s':d,'period_s':h+d,'conversion_percent':X,'paper_selectivity_percent':S,'ethane_per_inlet_A':X/100*(100/S-1),'cells':{'h':f'Figure 2 c!B{r}','d':f'Figure 2 c!C{r}','X':f'Figure 2 c!D{r}','S':f'Figure 2 d!D{r}'}}
 groups.setdefault(duty,[]).append(item)
for group in groups.values():group.sort(key=lambda x:x['period_s'])
out={'groups_with_exact_equal_duty':{k:v for k,v in groups.items() if len(v)>1},'uncertainty':'No replicate IDs, error columns or uncertainty bounds in Figure 2 b/c/d sheets. Conditional point-estimate comparisons only.'}
(p/'equal_duty_results.json').write_text(json.dumps(out,indent=2))
for k in ['10','100','1']:
 print(k,[(x['row'],x['ethane_per_inlet_A']) for x in groups[k]])
x,y=groups['10'][-2:];print('longest-pair',y['ethane_per_inlet_A']-x['ethane_per_inlet_A'],y['ethane_per_inlet_A']/x['ethane_per_inlet_A']-1)
x,y=groups['10'][0],groups['10'][-1];print('full-span',y['ethane_per_inlet_A']-x['ethane_per_inlet_A'],y['ethane_per_inlet_A']/x['ethane_per_inlet_A']-1)
