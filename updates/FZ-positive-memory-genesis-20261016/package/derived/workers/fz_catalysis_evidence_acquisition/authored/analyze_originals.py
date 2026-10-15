#!/usr/bin/env python3
"""Source-pinned reconstruction from publisher XLSX, not GC replication or experiment.

Run with the bundled Python runtime (openpyxl/numpy). Every source is opened read-only.
Original fulltexts and source arrays remain outside the authored publication directory.
"""
from pathlib import Path
import hashlib
import json
import math
import statistics
import openpyxl
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
EXT = ROOT / "external_excluded_from_publication"
OUT = ROOT / "authored"
FARADAY = 96485.33212
M_PD = 106.42
MASS_PD_G = 0.0001
PD_MOL = MASS_PD_G / M_PD
PINS = {
    "final_Fig2_MOESM20.xlsx": "47a571bed2750c89ba2cf38cc75235f1e7497727c57e2b9b28fc3fdfc667a3fe",
    "final_Fig6_MOESM24.xlsx": "28af31650f80fe130fef680035e127ab2f257ae681395d52ddb850662132fdfa",
    "final_SIFig7_MOESM8.xlsx": "c41ad6efa45d475db18b87ca2e88cbaf967f2e2fbfb31f8969a300c2c631d72f",
    "final_SIFig11_MOESM12.xlsx": "953c5f0a13da5588d1f853ef8b4c8bddb2d1cab8ec3f7ea9ed2d1f0465ddc983",
}


def workbook(name):
    path = EXT / name
    if name in PINS:
        observed = hashlib.sha256(path.read_bytes()).hexdigest()
        assert observed == PINS[name], (name, observed)
    return openpyxl.load_workbook(path, data_only=False, read_only=True)


def flux(x, s):
    b = x * (1 / s - 1)
    return {"paper_conversion": x, "paper_selectivity": s,
            "ethane_per_incoming_acetylene": b,
            "net_ethylene_per_incoming_acetylene_C2_only": x - b,
            "net_ethylene_per_consumed_acetylene_C2_only": 2 - 1 / s}


def contraction(x, s, z=0.01):
    # Literal mole-fraction scenario with exactly two hydrogenation steps.
    # Not established by source flow calibration or raw species balance.
    q = (1 - z) / (1 - z * (1 - x) + z * x * (1 / s - 1))
    true_x = 1 - q * (1 - x)
    b = q * x * (1 / s - 1)
    return {"outlet_to_inlet_total_molar_flow": q,
            "molar_flow_conversion": true_x,
            "ethane_per_incoming_acetylene": b,
            "net_ethylene_per_incoming_acetylene": true_x - b}


def bounds(static, dynamic, tp, tm, actual_feed_mol_s):
    bp = static["ethane_per_incoming_acetylene"]
    bd = dynamic["ethane_per_incoming_acetylene"]
    r = bd / bp
    avoided = actual_feed_mol_s * (tp * bp - (tp + tm) * bd)
    return {"positive_duration_s": tp, "negative_duration_s": tm,
            "dynamic_to_static_ethane_ratio": r,
            "memoryless_positive_endpoint_mixture_lower_bound": tp / (tp + tm),
            "positive_phase_ethane_mean_ratio_upper_bound": (tp + tm) / tp * r,
            "suppression_equivalent_seconds_not_lifetime": tp - (tp + tm) * r,
            "extra_avoided_ethane_mol_per_cycle_beyond_zero_negative_phase": avoided,
            "extra_avoided_events_per_total_Pd_atom": avoided / PD_MOL}


wb = workbook("final_Fig2_MOESM20.xlsx")
s = wb["Figure 2 b"]
endpoints = {
    "no_applied_voltage_row4": flux(s["C4"].value / 100, s["D4"].value / 100),
    "static_positive_row10": flux(s["C10"].value / 100, s["D10"].value / 100),
    "dynamic_10s_0p1s_row11": flux(s["C11"].value / 100, s["D11"].value / 100),
}
assert s["B11"].value == "1.5V 10s+-0.1V0.1s"
xc = {(r[1].value, r[2].value): (r[3].value / 100, r[3].coordinate) for r in wb["Figure 2 c"].iter_rows(min_row=2) if isinstance(r[1].value, (int, float))}
sd = {(r[1].value, r[2].value): (r[3].value / 100, r[3].coordinate) for r in wb["Figure 2 d"].iter_rows(min_row=2) if isinstance(r[1].value, (int, float))}
assert set(xc) == set(sd)
pulse_rows = {}
for tp, tm in [(10, 0.1), (10, 1)]:
    x, xloc = xc[tp, tm]
    se, sloc = sd[tp, tm]
    pulse_rows[f"{tp}s_{tm}s"] = {**flux(x, se), "conversion_cell": "Figure 2 c!" + xloc, "selectivity_cell": "Figure 2 d!" + sloc}
wb.close()
xp = endpoints["static_positive_row10"]
xd = endpoints["dynamic_10s_0p1s_row11"]
# Final SI Table1 scale and independent explicit volumetric-reference scenario.
activity = 9.918
feed_from_activity = activity * MASS_PD_G / 3600 / xd["paper_conversion"]
feed_298K_1atm = 0.0005 / (0.082057366080960 * 298.15) / 60
final_bound = bounds(xp, xd, 10, 0.1, feed_from_activity)
final_bound_feed = bounds(xp, xd, 10, 0.1, feed_298K_1atm)
corrected = {name: contraction(row["paper_conversion"], row["paper_selectivity"]) for name, row in endpoints.items()}
corrected_bound = bounds(corrected["static_positive_row10"], corrected["dynamic_10s_0p1s_row11"], 10, 0.1,
                         activity * MASS_PD_G / 3600 / corrected["dynamic_10s_0p1s_row11"]["molar_flow_conversion"])

wb = workbook("final_SIFig7_MOESM8.xlsx")
ws = wb["Figure 7 b"]
rows = np.array([[r[1].value, r[2].value, r[3].value] for r in ws.iter_rows(min_row=2)], dtype=float)
t, v, i = rows.T
dt = np.diff(t)
assert len(t) == 54999 and np.all(dt > 0) and np.max(np.abs(dt - 0.0002)) < 1e-10
assert ws["C1"].value == "Potential / V" and ws["D1"].value == "Current / A"
integral = getattr(np, "trapezoid", np.trapz)


def integ(y, lo=t[0], hi=t[-1]):
    take = (t >= lo) & (t <= hi)
    return float(integral(y[take], x=t[take]))


p = v * i
current = {
    "source": "final_SIFig7_MOESM8.xlsx:Figure 7 b!B2:D55000",
    "waveform_identity": "+1.5 V 10 s / -0.1 V 1 s, not optimum 0.1 s pulse",
    "samples": len(t), "first_time_s": float(t[0]), "last_time_s": float(t[-1]), "dt_s": float(np.median(dt)),
    "potential_range_V": [float(v.min()), float(v.max())],
    "current_range_A": [float(i.min()), float(i.max())],
    "minimum_current_value_occurrences": int(np.sum(i == i.min())),
    "maximum_current_value_occurrences": int(np.sum(i == i.max())),
    "charge_signed_C": integ(i), "charge_absolute_C": integ(abs(i)),
    "positive_current_charge_C": integ(np.maximum(i, 0)),
    "negative_current_charge_magnitude_C": integ(np.maximum(-i, 0)),
    "recorded_terminal_VI_signed_J": integ(p),
    "recorded_terminal_VI_absolute_J": integ(abs(p)),
    "recorded_terminal_VI_positive_J": integ(np.maximum(p, 0)),
    "recorded_terminal_VI_negative_magnitude_J": integ(np.maximum(-p, 0)),
    "recorded_terminal_VI_mean_W_over_available_span": integ(p) / (t[-1] - t[0]),
    "integral0p0002_to1s": {"signed_charge_C": integ(i, t[0], 1), "signed_VI_J": integ(p, t[0], 1)},
    "integral1s_to10p9998s": {"signed_charge_C": integ(i, 1, t[-1]), "signed_VI_J": integ(p, 1, t[-1])},
    "positive_tail_average_current_mA_by_start_s": {str(start): integ(i, start, t[-1]) / (t[-1] - start) * 1000 for start in (1.5, 2, 3, 4, 5)},
    "positive_tail_VI_J_from2s": integ(p, 2, t[-1]),
    "charge_absolute_electron_equivalents_per_total_Pd_atom": integ(abs(i)) / FARADAY / PD_MOL,
    "qualifications": [
        "Trapezoidal integration of published numeric waveform; no independent instrument calibration.",
        "First and last 0.2 ms endpoint gaps are not imputed.",
        "Repeated extreme currents are reported; possible range saturation cannot be excluded without instrument log.",
        "VI is the integral of recorded terminal channels, not a wall-plug or total-process energy measurement.",
        "Charge is not identified with hydrogen uptake; charging and Faradaic components are not separable from these columns alone.",
        "No 10s/0.1s current waveform is supplied here; this energy must not be assigned to optimum performance."
    ]
}
current_grid = {}
for wsname in ("Figure 7 c", "Figure 7 d"):
    ws = wb[wsname]
    current_grid[wsname] = [{"tpositive_s": r[1].value, "tnegative_s": r[2].value, "value_as_stored": r[3].value, "cell": r[3].coordinate}
                            for r in ws.iter_rows(min_row=2) if (r[1].value, r[2].value) in {(10, .1), (10, 1)}]
wb.close()

wb = workbook("final_Fig6_MOESM24.xlsx")
spectra = []
for ws in wb:
    vals = list(ws.values)
    time_header = next((row for row in vals[:3] if "t / s" in row), None)
    times = [x for x in time_header if isinstance(x, (int, float))]
    time_columns = [j for j, x in enumerate(time_header) if isinstance(x, (int, float))]
    first_data_row = next(j for j, row in enumerate(vals) if isinstance(row[time_columns[0] - 1], (int, float)) and j > 1)
    axis_col = time_columns[0] - 1
    axis_missing_rows = [j + 1 for j, row in enumerate(vals) if j >= first_data_row and
                         not isinstance(row[axis_col], (int, float)) and any(isinstance(row[col], (int, float)) for col in time_columns)]
    values = [row for row in vals[first_data_row:] if isinstance(row[axis_col], (int, float))]
    arrays = [tuple(row[col] for row in values) for col in time_columns]
    duplicate_adjacent = [[times[j], times[j+1]] for j in range(len(arrays)-1) if arrays[j] == arrays[j+1]]
    spectra.append({"sheet": ws.title, "spectra": len(times), "first_time_s": min(times), "last_time_s": max(times),
                    "time_spacing_s": sorted(set(np.diff(times).tolist())), "wavenumber_points": len(values),
                    "wavenumber_range_cm_1": [values[0][axis_col], values[-1][axis_col]],
                    "source_rows_with_spectrum_values_but_missing_wavenumber": axis_missing_rows,
                    "exactly_identical_adjacent_spectra_times_s": duplicate_adjacent,
                    "formula_cells": sum(isinstance(x, str) and x.startswith("=") for row in vals for x in row)})
wb.close()

wb = workbook("final_SIFig11_MOESM12.xlsx")
control = wb["Figure 11 c"]
control12h = {"source": "Figure 11 c!B5:D5", "condition": control["B5"].value,
              "column_C_C2H4_to_C2H6_mol_min": control["C5"].value,
              "column_D_C2H2_to_C2H4_mol_min": control["D5"].value}
wb.close()

results = {
    "status": "original_publisher_sources_acquired_and_recalculated_no_new_experiment",
    "source_sha256": PINS,
    "final_Fig2_endpoints": endpoints,
    "final_Fig2_pulse_grid_rows": pulse_rows,
    "final_bound_activity_scale": final_bound,
    "final_bound_nominal_feed_298p15K_1atm_scenario": final_bound_feed,
    "final_ideal_two_step_flow_contraction_scenario": {"rows": corrected, "bounds_activity_scale": corrected_bound},
    "scale": {"total_Pd_mol_nominal": PD_MOL, "feed_acetylene_mol_s_from_final_SI_activity": feed_from_activity,
              "nominal_0p5mL_min_at298p15K_1atm_feed_mol_s": feed_298K_1atm,
              "final_SI_activity_mol_gPd_h": activity,
              "net_ethylene_productivity_mol_gPd_h_C2_only": activity * xd["net_ethylene_per_consumed_acetylene_C2_only"],
              "final_SI_reported_productivity_mol_gPd_h": 7.525},
    "comparison_to_supplied_producer_rounding": {
        "source": "supplied results.json:conditional_dynamic_test, exact scalar transcription",
        "ethane_ratio_producer": .2214772280049756,
        "ethane_ratio_final_minus_producer": final_bound["dynamic_to_static_ethane_ratio"] - .2214772280049756,
        "suppression_equivalent_seconds_producer": 7.763079997149747,
        "suppression_equivalent_seconds_final_minus_producer": final_bound["suppression_equivalent_seconds_not_lifetime"] - 7.763079997149747,
        "events_per_Pd_producer": 2.489364953316058,
        "events_per_Pd_final_minus_producer": final_bound["extra_avoided_events_per_total_Pd_atom"] - 2.489364953316058},
    "complete_published_current_trace_10s_1s": current,
    "SI7_current_grid_cells_unit_as_stored_unresolved": current_grid,
    "Fig6_DRIFTS_source_structure": spectra,
    "SI11_low_current_12h_control_source": control12h,
    "assumptions_for_suppression_bound": [
        "Common flow-normalized molar-flux interpretation of conversion and selectivity, or explicit contraction sensitivity.",
        "Shared reactor/specimen-class/calibration and reproducible steady periodic endpoints.",
        "Nonnegative phase-resolved ethane production and closed integrated storage over a period.",
        "C2-only reconstructed net ethylene; side products and carbon deficits require separate treatment.",
        "The one-purge-per-Pd exclusion is limited to total metal sites; support and pore inventories were not acquired."
    ],
    "unacquired_or_unestablished": [
        "Raw GC chromatograms, peak integrals, standards and response-factor calibration records.",
        "Measured inlet/outlet total molar flows or nitrogen-normalized species fluxes.",
        "Independent preparation/replicate uncertainties and pulse-order/run-history records.",
        "Species-resolved output trace across the 0.1s-pulse/10s-hold cycle.",
        "Matched-system accessible inventory and hold-only recovery measurements.",
        "Complete optimum-waveform electrical integral and wall-plug/process energy.",
        "Final main-paper fulltext binary; preprint v1 and final SI/figure sources are version-separated."
    ]
}
assert 0.21 < final_bound["dynamic_to_static_ethane_ratio"] < .23
assert 7.7 < final_bound["suppression_equivalent_seconds_not_lifetime"] < 7.8
assert 2.4 < final_bound["extra_avoided_events_per_total_Pd_atom"] < 2.6
assert math.isclose(pulse_rows["10s_0.1s"]["paper_conversion"], xd["paper_conversion"], abs_tol=.001)
assert len(xc) == 30
(OUT / "analysis_results.json").write_text(json.dumps(results, indent=2) + "\n")
print(json.dumps({"final_bound_activity_scale": final_bound, "pulse_rows": pulse_rows, "current": current,
                  "current_grid_cells": current_grid, "spectra_structure": spectra, "low_current_control": control12h}, indent=2))
