"""Read-only audit of original source workbooks with bundled Python libraries.

Never overwrite, impute, or repair a source cell. Crossing calculations describe
the stored normalized SQUID curves; they do not certify field calibration.
"""
from pathlib import Path
import datetime
import json
import numbers
import numpy as np
from openpyxl import load_workbook

B=Path(__file__).resolve().parent
P=B/'data/raw_loops'
w=load_workbook(P/'Magnetizm_FeCoNiCu_NWs.xlsx',read_only=True,data_only=True)
s=w['hysteresis_loops'];rows=list(s.values)
parameters={float(r[0]):{'source_Hc_Cu_Oe':r[1],'source_Hc_Cufree_Oe':r[3],
                         'source_squareness_Cu_values':r[2], 'source_squareness_Cufree_values':r[4]}
            for r in list(w['magnetic_parameters'].values)[2:]}
loops=[]
for col in range(0,16,2):
    points=[(row[col],row[col+1],i+1) for i,row in enumerate(rows[2:],start=2)
            if isinstance(row[col],numbers.Real) and isinstance(row[col+1],numbers.Real)]
    z=np.array([[h,m] for h,m,row in points]);h=z[:,0];m=z[:,1]
    crossing=[]
    for i in range(len(m)-1):
        if m[i]*m[i+1]<0:
            crossing.append({'linear_crossing_Oe':float(h[i]-m[i]*(h[i+1]-h[i])/(m[i+1]-m[i])),
                             'field_bracket_Oe':sorted([float(h[i]),float(h[i+1])]),
                             'source_rows_1based':[points[i][2],points[i+1][2]],
                             'M_over_Ms_bracket':[float(m[i]),float(m[i+1])]})
    assert len(crossing)==2
    crossing.sort(key=lambda d:d['linear_crossing_Oe'])
    neg,pos=crossing;hc=.5*(pos['linear_crossing_Oe']-neg['linear_crossing_Oe'])
    bias=.5*(pos['linear_crossing_Oe']+neg['linear_crossing_Oe'])
    interval=[.5*(pos['field_bracket_Oe'][0]-neg['field_bracket_Oe'][1]),
              .5*(pos['field_bracket_Oe'][1]-neg['field_bracket_Oe'][0])]
    voltage=float(rows[1][col+1].split()[0]);geometry='out_of_membrane_plane' if col<8 else 'in_membrane_plane'
    record={'potential_V':voltage,'geometry':geometry,'n_points':len(points),
            'source_sheet':'hysteresis_loops','source_columns_1based':[col+1,col+2],
            'linear_interpolation_Hc_Oe':float(hc),'crossing_bracket_Hc_interval_Oe':interval,
            'loop_center_Oe':float(bias),'crossings':crossing,
            'field_and_normalized_magnetization':z.tolist()}
    if col<8:
        tabulated=parameters[voltage]['source_Hc_Cu_Oe']
        record['source_tabulated_Hc_Oe']=tabulated
        record['table_minus_linear_interpolation_Oe']=float(tabulated-hc)
        record['tabulated_value_inside_crossing_bracket_interval']=interval[0]<=tabulated<=interval[1]
    loops.append(record)

wx=load_workbook(P/'XRD_FeCoNiCu_NWs.xlsx',read_only=True,data_only=True)
sx=wx['XRD_patterns'];xr=list(sx.values);date_cells=[];bad_numeric=[];valid_counts=[]
for col in [0,2,4,6]:
    nvalid=0
    for i,row in enumerate(xr[2:],start=3):
        value=row[col]
        if isinstance(value,(datetime.date,datetime.datetime)):
            date_cells.append({'row_1based':i,'column_1based':col+1,'value':value.isoformat()})
        elif isinstance(value,numbers.Real) and 0<=value<=180:
            nvalid+=1
        elif value is not None:
            bad_numeric.append({'row_1based':i,'column_1based':col+1,'value':str(value)})
    valid_counts.append({'column_1based':col+1,'valid_numeric_2theta_points':nvalid})

angular=[]
for i,row in enumerate(list(w['Angular_parameters'].values)[2:],start=3):
    if isinstance(row[0],numbers.Real):angular.append({'source_row_1based':i,'angle_deg':row[0],
                                                      'values':[row[1],row[2],row[3]]})

out={'scope':'audit_of_actual_source_normalized_SQUID_curves_not_independent_instrument_truth',
     'source_doi':'10.48733/IFJPAN/VV0KPY','original_workbook':'data/raw_loops/Magnetizm_FeCoNiCu_NWs.xlsx',
     'loops':loops,'source_parameters_by_potential':parameters,
     'data_quality':{
       'Dataverse_tab_export_omits_five_magnetic_worksheets':True,
       'XRD_dates_in_original_workbook':date_cells,
       'XRD_invalid_numeric_angle_cells':bad_numeric,
       'XRD_valid_numeric_counts':valid_counts,
       'angular_header':'Coercivity [Oe]',
       'angular_values_are_one_at_zero_angle':'All three series equal 1 at angle=0; therefore header alone cannot justify treating values as absolute Oe.',
       'angular_table_records':angular,
       'angular_suspicious_cell':{'sheet':'Angular_parameters','cell':'B11','value':475548,'status':'unresolved; preserve original and do not impute'},
       'source_loop_vs_parameter_mismatch':[{'potential_V':r['potential_V'],
                    'tabulated_Oe':r['source_tabulated_Hc_Oe'],'loop_linear_interpolation_Oe':r['linear_interpolation_Hc_Oe'],
                    'inside_crossing_bracket':r['tabulated_value_inside_crossing_bracket_interval']}
                    for r in loops if r['geometry']=='out_of_membrane_plane'],
       'independent_field_calibration_standard':None,'repeatability_error_bars':None,
       'absolute_Ms_from_stored_normalized_loops':None,
       'all_eight_interpolated_crossings_are_exactly_symmetric':all(r['loop_center_Oe']==0 for r in loops),
       'symmetry_qualification':'Stored opposite branches are not evidence of independent replicate measurements. Processing and symmetrization provenance remain unresolved.'},
     'interpretation':'Public raw-loop access partially repairs functional access. Direction and processing change the target. Missing calibration and table discrepancies prevent precision or transfer claims. Four deposition conditions are a validation audit, not a credible large-scale search benchmark.',
     'new_operation_demonstrated':False}
(B/'RAW_LOOP_AUDIT.json').write_text(json.dumps(out,indent=2)+'\n')
print('LOOPS')
for r in loops:print(r['potential_V'],r['geometry'],r['linear_interpolation_Hc_Oe'],r['crossing_bracket_Hc_interval_Oe'],r.get('source_tabulated_Hc_Oe'),r.get('tabulated_value_inside_crossing_bracket_interval'))
print('XRD date cells',len(date_cells),'invalid numeric cells',len(bad_numeric));print('angular suspicious B11',w['Angular_parameters']['B11'].value)
