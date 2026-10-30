#!/usr/bin/env python3
"""Read-only source audit. Run with bundled Python containing numpy/openpyxl.

No workbook is modified. Source labels are kept distinct from measured traces,
model constructions, and physical performance claims. No network is needed.
"""
from __future__ import annotations
import hashlib
import itertools
import json
import math
from pathlib import Path
import re
import time

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "sources"
F = 96485.0
M_CO = 0.02801  # kg/mol, used only for a derived electrode-power proxy


def mean_time(x, t):
    return float(np.trapezoid(x, t) / (t[-1] - t[0]))


def provenance():
    receipts = []
    for meta, prefix, persistent_id in [
        ("edmond_metadata.json", "", "doi:10.17617/3.CCHCNR"),
        ("edmond_power2025_metadata.json", "power2025_", "doi:10.17617/3.FEEHJO"),
    ]:
        d = json.loads((SRC / meta).read_text())["data"]
        ver = d["latestVersion"]
        for x in ver["files"]:
            f = x["dataFile"]
            p = SRC / (prefix + f["filename"])
            if not p.exists():
                continue
            b = p.read_bytes()
            expected = f["checksum"]["value"]
            actual = hashlib.md5(b).hexdigest()
            assert actual == expected, p
            receipts.append({
                "file": str(p.relative_to(ROOT)), "bytes": len(b),
                "dataset": persistent_id,
                "dataset_version": [ver.get("versionNumber"), ver.get("versionMinorNumber")],
                "url": f"https://edmond.mpg.de/api/access/datafile/{f['id']}",
                "publisher_md5": expected, "verified_md5": actual,
                "sha256": hashlib.sha256(b).hexdigest(),
            })
    for p in sorted(SRC.glob("*.pdf")):
        b = p.read_bytes()
        receipts.append({"file": str(p.relative_to(ROOT)), "bytes": len(b),
                         "sha256": hashlib.sha256(b).hexdigest(),
                         "publisher_checksum_available": False})
    return receipts


def rde_audit():
    p = next(SRC.rglob("figure_3_raw_data_calculated_edmond.xlsx"))
    ws = openpyxl.load_workbook(p, data_only=True, read_only=True).active
    points = []
    for row in range(2, 8):
        v = [ws.cell(row, c).value for c in range(1, 10)]
        freq, amp, pot, curr, power, ppm, ico, fe_label, fe_calc = v
        points.append({
            "frequency_Hz": freq, "amplitude_RMS_V": amp,
            "nominal_potential_RHE_V": pot, "mean_current_A": curr,
            "source_average_power_W": power, "CO_ppm": ppm,
            "CO_equivalent_current_A": -ico, "source_FE_label": fe_label,
            "recomputed_FE_from_source_current": -ico / -curr,
            "source_power_per_CO_current_V": power / -ico,
            "source_power_proxy_kWh_kgCO": power * 2 * F / (-ico * M_CO * 3.6e6),
            "source_cell": f"Sheet1!A{row}:I{row}",
        })
    # Exact affine reference dependence of any single-electrode energy proxy.
    # c adds to cathodic voltage magnitude; this is a diagnostic, not a cell model.
    a, b = points[1], points[5]
    slope_a = -a["mean_current_A"] / a["CO_equivalent_current_A"]
    slope_b = -b["mean_current_A"] / b["CO_equivalent_current_A"]
    intercept_a = a["source_power_per_CO_current_V"]
    intercept_b = b["source_power_per_CO_current_V"]
    crossover = -(intercept_a - intercept_b) / (slope_a - slope_b)
    witness = {
        "protocol_A": [a["frequency_Hz"], a["amplitude_RMS_V"]],
        "protocol_B": [b["frequency_Hz"], b["amplitude_RMS_V"]],
        "definition": "g(c)=(P_source+c*abs(mean_I))/I_CO; conditional single-electrode proxy diagnostic",
        "intercepts_V": [intercept_a, intercept_b],
        "slopes": [slope_a, slope_b],
        "crossover_reference_offset_V": crossover,
        "values_at_offset_0": [intercept_a, intercept_b],
        "values_at_offset_0_6562": [intercept_a + .6562 * slope_a,
                                     intercept_b + .6562 * slope_b],
        "scope": "Neither ordering is a ranking of full cell energy; full terminal voltage is missing.",
    }
    traces = []
    for p in sorted(SRC.rglob("*.xlsx")):
        if "/dynamic/" not in str(p):
            continue
        ws = openpyxl.load_workbook(p, data_only=False, read_only=True).active
        reduced = ws.max_column == 3
        formula = ws["C5" if reduced else "D5"].value
        lo, hi = map(int, re.findall(r"\d+", formula))
        ic = 2 if reduced else 3
        a = np.array(list(ws.iter_rows(min_row=lo, max_row=hi,
                                      max_col=ic, values_only=True)), dtype=float)
        t, current = a[:, 0], a[:, ic-1]
        freq = float(re.search(r"_(\d+(?:\.\d+)?)Hz", p.name)[1])
        amp = float(re.search(r"amprms_(\d+\.\d+)V", p.name)[1])
        source_point = next(z for z in points if z["frequency_Hz"] == freq
                            and z["amplitude_RMS_V"] == amp)
        current_difference = float(current.mean() - source_point["mean_current_A"])
        # The 1 Hz, 0.3 V summary rounds its current to 9 significant figures.
        assert abs(current_difference) < 2e-10
        item = {
            "file": str(p.relative_to(ROOT)), "sheet": ws.title,
            "source_average_formula": formula, "window_rows_inclusive": [lo, hi],
            "n_samples": len(current), "mean_current_A": float(current.mean()),
            "source_current_replay_difference_A": current_difference,
            "sample_mean_current_squared_A2": float(np.mean(current**2)),
            "time_span_s": [float(t[0]), float(t[-1])],
            "dt_min_median_max_s": np.quantile(np.diff(t), [0, .5, 1]).tolist(),
            "positive_current_sample_count": int(np.sum(current > 0)),
            "voltage_column_available": not reduced,
            "header_current_units": str(ws.cell(4, ic).value),
        }
        if not reduced:
            e = a[:, 1]
            # Source conversion E_RHE = E_AgAgCl + .059*pH + .196, pH 7.8.
            raw_proxy = float(np.mean((e + .059 * 7.8 + .196) * current))
            item.update({"replayed_mean_electrode_product_RHE_W": raw_proxy,
                         "corrected_table_power_W": source_point["source_average_power_W"],
                         "replayed_minus_corrected_table_W": raw_proxy-source_point["source_average_power_W"],
                         "discrepancy_status": "UNRESOLVED: source averaging/correction protocol for V2 power not supplied in workbook"})
        else:
            item["power_replay_status"] = "UNAVAILABLE: current-only release; irregular samples much slower than 100 Hz period; no acquisition/reduction rule"
        traces.append(item)
    return {"corrected_source_file":
            str(next(SRC.rglob("figure_3_raw_data_calculated_edmond.xlsx")).relative_to(ROOT)),
            "corrected_points": points, "reference_dependence_witness": witness,
            "trace_replays": traces}


def flow_audit():
    static_book = openpyxl.load_workbook(next(SRC.rglob("*Steady.xlsx")),
                                       data_only=True, read_only=True)
    dynamic_book = openpyxl.load_workbook(next(SRC.rglob("*Dynamic.xlsx")),
                                        data_only=True, read_only=True)
    dynamics = []
    for ws in dynamic_book:
        for row in (3, 4):
            r = [ws.cell(row, c).value for c in range(1, 12)]
            ssrow = row + 2
            ssq, ssi = -ws.cell(ssrow, 9).value, -ws.cell(ssrow, 10).value
            dynamics.append({"sheet": ws.title, "row": row,
                "annotation": r[10], "CO_equivalent_current_A": -r[8],
                "mean_cathodic_current_A": -r[9], "FE": r[8]/r[9],
                "same_nominal_potential_static_CO_current_A": ssq,
                "same_nominal_potential_static_FE": ssq/ssi,
                "CO_gain_vs_same_nominal_potential": (-r[8])/ssq,
                "H2_dynamic_assay_available": False,
                "full_cell_terminal_voltage_available": False})
    mixes = []
    for ws in static_book:
        pts = [(r[0], -r[10], -r[11]) for r in ws.iter_rows(
            min_row=3, max_row=11, max_col=12, values_only=True)
            if isinstance(r[0], (int, float))]
        for d in dynamics:
            q = d["CO_equivalent_current_A"]
            candidates = []
            for a, b in itertools.combinations(pts, 2):
                if a[1] == b[1]:
                    continue
                weight_b = (q-a[1])/(b[1]-a[1])
                if 0 <= weight_b <= 1:
                    j = (1-weight_b)*a[2]+weight_b*b[2]
                    candidates.append((j,a,b,weight_b))
            best = min(candidates)
            mixes.append({"static_batch": ws.title,
                "dynamic_sheet": d["sheet"], "dynamic_row": d["row"],
                "CO_target_A": q, "static_minimum_mixture_current_A": best[0],
                "pulse_current_A": d["mean_cathodic_current_A"],
                "static_to_pulse_current_ratio": best[0]/d["mean_cathodic_current_A"],
                "static_modes_EAgCl_V_Ico_A_I_A": [best[1],best[2]],
                "weight_on_second_mode": best[3],
                "status": "HISTORICAL_DESCRIPTIVE_ONLY: different batch/date/flow; H2, total power, transition and peak constraints unmatched"})
    return {"measured_source_labels": dynamics, "historical_charge_envelopes": mixes}


def full_cell_incumbent():
    r = SRC / "power2025_experiments" / "Experimental results"
    ss = np.loadtxt(r / "exp_data_power_1000_mA_SS.txt", delimiter=",", skiprows=1)[0]
    summaries = []
    for amp in (10, 25, 50):
        a = np.loadtxt(r/f"exp_data_power_1000_mA_{amp}_percent.txt", delimiter=",", skiprows=1)
        delta = 100*(a[:, 1]/ss[1]-1)
        summaries.append({"amplitude_percent_label": amp, "frequency_count": len(a),
            "frequency_range_Hz": [float(a[0,0]),float(a[-1,0])],
            "relative_mean_power_difference_percent_min_max": [float(delta.min()),float(delta.max())],
            "means_above_static_count": int(np.sum(a[:,1]>ss[1])),
            "means_below_static_count": int(np.sum(a[:,1]<ss[1])),
            "source_replicate_count": 3,
            "interpretation": "Signs of source means only; not independent significance tests or gas-yield measurements."})
    traces = []
    for p in sorted(r.glob("Voltage_current_signals_cut*")):
        a = np.loadtxt(p, skiprows=1)
        t, i, v = a.T
        mi, mv, power = mean_time(i,t), mean_time(v,t), mean_time(i*v,t)
        traces.append({"file": str(p.relative_to(ROOT)), "n_samples": len(i),
            "time_span_s": [float(t[0]),float(t[-1])],
            "mean_current_density_A_cm2": mi, "mean_cell_voltage_V": mv,
            "mean_power_density_W_cm2": power, "product_of_means_W_cm2": mi*mv,
            "covariance_W_cm2": power-mi*mv,
            "mean_product_underestimate_percent": 100*(power/(mi*mv)-1),
            "charge_normalized_V": power/mi,
            "source_static_power_density_W_cm2": float(ss[1]),
            "pointwise_power_difference_vs_source_static_percent": 100*(power/ss[1]-1),
            "integration": "trapezoidal time average over source cut interval; repeated endpoint not counted twice"})
    return {"dataset": "doi:10.17617/3.FEEHJO", "static_power_density_W_cm2":float(ss[1]),
            "static_source_standard_deviation_W_cm2":float(ss[2]),
            "scope": "Measured full-cell power; fixed mean current. Equal H2 output uses source 100-percent-Faradaic-efficiency assumption, not acquired independent gas-output evidence.",
            "frequency_summaries":summaries,"pointwise_trace_replays":traces}


def identity_check():
    # Algebraic verification of the theorem, not an experimental model fit.
    t = np.arange(2**17)*(2*math.pi/2**17)
    x = .7 + .2*np.sin(t) + .05*np.sin(3*t)
    xdot = .2*np.cos(t) + .15*np.cos(3*t)
    f = x + x**3
    c = 1 + x**2
    resistance = .4
    i = f + c*xdot
    u = x + resistance*i
    lhs = float(np.mean(u*i))
    static = float(np.mean(x*f + resistance*f*f))
    excess = float(resistance*np.mean((c*xdot)**2))
    assert abs(lhs-static-excess)<1e-12
    return {"kind":"MATHEMATICAL_CHECK_ONLY", "mean_input_power":lhs,
            "static_occupation_power":static,"nonnegative_ripple_term":excess,
            "identity_residual":lhs-static-excess,
            "statement":"An arbitrary-amplitude periodic potential in the specified one-state capacitor model has the same product occupation means as static time-sharing and an extra nonnegative resistor cost."}


def main():
    started = time.perf_counter()
    out = {"status":"HOLD_NO_NEW_CAPABILITY", "worker":"a05_energy_matter",
           "source_receipts":provenance(), "co2_rde":rde_audit(),
           "co2_flow":flow_audit(), "full_cell_incumbent":full_cell_incumbent(),
           "occupation_identity_check":identity_check()}
    out["local_replay_wall_seconds"] = time.perf_counter()-started
    out["inference_tokens"] = None
    out["inference_cost"] = None
    out["inference_energy"] = None
    (ROOT / "RESULTS.json").write_text(json.dumps(out,indent=2,ensure_ascii=False)+"\n")
    print(json.dumps({"result":"RESULTS.json", "status":out["status"],
                      "source_receipt_count":len(out["source_receipts"]),
                      "replayed_rde_current_windows":len(out["co2_rde"]["trace_replays"]),
                      "full_cell_waveform_replays":len(out["full_cell_incumbent"]["pointwise_trace_replays"]),
                      "full_cell_frequency_summary_count":sum(x["frequency_count"] for x in out["full_cell_incumbent"]["frequency_summaries"]),
                      "local_replay_wall_seconds":out["local_replay_wall_seconds"]},indent=2))


if __name__ == "__main__":
    main()
