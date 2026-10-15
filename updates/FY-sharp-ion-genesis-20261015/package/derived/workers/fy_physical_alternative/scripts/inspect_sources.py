"""Read-only extraction; does not rewrite source PDF/XLSX bytes."""
from pathlib import Path
import json
from datetime import datetime, date
import openpyxl
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "derived"
OUT.mkdir(exist_ok=True)
index = []
for path in sorted((ROOT / "sources").glob("*.xlsx")):
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    tabs = []
    for tab in book.worksheets:
        values = []
        for row in tab.iter_rows():
            for cell in row:
                value = cell.value
                if value is not None:
                    if isinstance(value, (datetime, date)):
                        value = value.isoformat()
                    values.append({"cell": cell.coordinate, "value": value})
        tabs.append({"sheet": tab.title, "max_row": tab.max_row,
                     "max_column": tab.max_column, "cells": values})
        if tab.max_row < 20:
            print(path.name, tab.title, json.dumps(values), flush=True)
    (OUT / (path.stem + ".json")).write_text(json.dumps(tabs, indent=2) + "\n")
    index.append({"workbook": path.name, "tabs": [
        {key: tab[key] for key in ["sheet", "max_row", "max_column"]} for tab in tabs
    ]})
    book.close()
(OUT / "WORKBOOK_INDEX.json").write_text(json.dumps(index, indent=2) + "\n")
for path in sorted((ROOT / "sources").glob("*.pdf")):
    reader = PdfReader(path)
    pages = [f"\n--- PDF PAGE {i} ---\n" + (page.extract_text() or "")
             for i, page in enumerate(reader.pages, 1)]
    (OUT / (path.stem + ".txt")).write_text("\n".join(pages))
    print(path.name, "PDF pages", len(reader.pages), flush=True)
