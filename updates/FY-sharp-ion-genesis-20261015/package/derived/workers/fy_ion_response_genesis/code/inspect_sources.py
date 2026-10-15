"""Read-only source inspection. No source workbook or PDF is modified."""
from pathlib import Path
from html.parser import HTMLParser
import json
from openpyxl import load_workbook
from pypdf import PdfReader

BASE = Path(__file__).resolve().parents[1]


class ArticleText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.article_depth = 0
        self.blocked_depth = 0
        self.out = []

    def handle_starttag(self, tag, attrs):
        if tag == "article":
            self.article_depth += 1
        if tag in ("script", "style"):
            self.blocked_depth += 1

    def handle_endtag(self, tag):
        if tag == "article":
            self.article_depth -= 1
        if tag in ("script", "style"):
            self.blocked_depth -= 1

    def handle_data(self, data):
        if self.article_depth > 0 and self.blocked_depth == 0 and data.strip():
            self.out.append(data.strip())


for f in sorted((BASE / "third_party").glob("*.xlsx")):
    workbook = load_workbook(f, read_only=True, data_only=True)
    info = []
    print("FILE", f.name)
    for sheet in workbook:
        rows = []
        for row in sheet.iter_rows():
            nonempty = {c.coordinate: c.value for c in row if c.value is not None}
            if nonempty:
                rows.append(nonempty)
            if len(rows) >= 5:
                break
        entry = dict(sheet=sheet.title, rows=sheet.max_row,
                     cols=sheet.max_column, first_nonempty=rows)
        info.append(entry)
        print(dict(sheet=sheet.title, rows=sheet.max_row,
                   cols=sheet.max_column,
                   header_preview=list(rows[0].items())[:8] if rows else []))
    (BASE / "evidence" / f"{f.stem}_inventory.json").write_text(
        json.dumps(info, indent=2, default=str) + "\n")
    workbook.close()

for f in sorted((BASE / "third_party").glob("*.pdf")):
    reader = PdfReader(f)
    text = "\n".join(f"\n--- PDF PAGE {i+1} ---\n" +
                     (page.extract_text() or "")
                     for i, page in enumerate(reader.pages))
    (BASE / "evidence" / f"{f.stem}.txt").write_text(text)
    print("PDF", f.name, len(reader.pages), len(text))

for name in ["kautz_publisher", "chen_ed_publisher", "grossman_arxiv_v1"]:
    parser = ArticleText()
    parser.feed((BASE / "third_party" / f"{name}.html").read_text())
    text = "\n".join(parser.out)
    (BASE / "evidence" / f"{name}.txt").write_text(text)
    print("HTML", name, len(text))
    if name == "kautz_publisher":
        start = text.find("Source Data")
        print(text[start:start + 1800])
