"""Re-analyze published source values; no invented species or work data."""
from pathlib import Path
import csv
import json
import math
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1]
source = BASE / "third_party" / "kautz_source_data_MOESM4.xlsx"
wb = load_workbook(source, read_only=True, data_only=True)
sheet = wb["Fig4b"]
ed, ref, voltage = [], [], []
for row_number, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
    for output, x, y in [(ed, row[0], row[1]), (ref, row[2], row[3]),
                         (voltage, row[4], row[5])]:
        if (isinstance(x, (int, float)) and isinstance(y, (int, float))
                and math.isfinite(x) and math.isfinite(y)):
            output.append(dict(time_hours=x, value=y, source_row=row_number))
wb.close()
assert [x["time_hours"] for x in ed] == [x["time_hours"] for x in ref]
assert all(b["time_hours"] > a["time_hours"] for a, b in zip(ed, ed[1:]))
minimum = min(ed, key=lambda x: x["value"])
crossing_index = next(i for i, x in enumerate(ed) if x["value"] <= 0.5)
crossing_bracket = ed[crossing_index - 1:crossing_index + 1]
ref_at_min = next(x for x in ref if x["source_row"] == minimum["source_row"])

initial_moles = 1e-3 * 0.63e-3  # stated KCl concentration M and dilution volume L
inferred_removed = initial_moles * (1 - minimum["value"])
inferred_rate = inferred_removed / minimum["time_hours"]
inferred_area_flux = inferred_rate / (0.16e-4)
results = dict(
    source_file=source.name, source_sheet="Fig4b",
    measured_scope="Published normalized conductance and voltage; not species-resolved or new measurements",
    ed_points=len(ed), reference_points=len(ref), voltage_points=len(voltage),
    minimum=minimum, reference_at_minimum=ref_at_min,
    first_half_crossing_bracket=crossing_bracket,
    last_ed=ed[-1], last_reference=ref[-1],
    implied_KCl_removed_mol=inferred_removed,
    implied_KCl_rate_mol_per_elapsed_hour=inferred_rate,
    implied_KCl_flux_mol_per_m2_per_elapsed_hour=inferred_area_flux,
    conversion_assumptions=["dilute KCl conductance proportional to concentration",
                            "constant 0.63 mL dilution-cell volume",
                            "initial 1 mM KCl", "stated active area 0.16 cm2"],
    unresolved=["pump-on duty exposure not reconstructed from voltage threshold",
                "Li, Na and Cl species inventories absent",
                "no direct all-path work time series in this workbook",
                "no reset/electrode/membrane inventory assay",
                "no replicate-based uncertainty from this single series"],
)
(BASE / "evidence" / "conductance_analysis.json").write_text(
    json.dumps(results, indent=2) + "\n")
for name, values in [("ed", ed), ("reference", ref), ("voltage", voltage)]:
    with (BASE / "evidence" / f"published_{name}_values.csv").open("w") as stream:
        writer = csv.DictWriter(stream, fieldnames=["time_hours", "value", "source_row"])
        writer.writeheader()
        writer.writerows(values)

print(json.dumps(results, indent=2))
