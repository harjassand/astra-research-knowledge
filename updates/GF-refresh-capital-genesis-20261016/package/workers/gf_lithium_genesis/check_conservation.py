#!/usr/bin/env python3
"""Exact conservation checks and read-only provenance replay, not a chemistry model."""
from pathlib import Path
from itertools import product
import hashlib
import json
import openpyxl

BASE = Path(__file__).resolve().parent
SOURCE = BASE.parents[1] / "originals/expanded/lithium_control_handoff/lithium_control_handoff"

def isotope_interval(a, b, z):
    lower, upper = max(0, z-b), min(z, a)
    if lower > upper:
        raise ValueError("Inconsistent isotope pool/input budgets")
    return lower, upper

cases = []
for a, b in product(range(7), repeat=2):
    for z in range(a+b+1):
        feasible = [f for f in range(a+1) if 0 <= z-f <= b]
        interval = isotope_interval(a, b, z)
        assert feasible
        assert interval == (min(feasible), max(feasible))
        cases.append((a, b, z, feasible))

paired_cases = 0
for left, right in product(cases, repeat=2):
    feasible_totals = {x+y for x in left[3] for y in right[3]}
    l0, u0 = isotope_interval(*left[:3])
    l1, u1 = isotope_interval(*right[:3])
    assert min(feasible_totals) == l0+l1
    assert max(feasible_totals) == u0+u1
    assert feasible_totals == set(range(l0+l1, u0+u1+1))
    paired_cases += 1

inconsistent = 0
for a, b in product(range(7), repeat=2):
    try:
        isotope_interval(a, b, a+b+1)
    except ValueError:
        inconsistent += 1
    else:
        raise AssertionError("Missing inconsistent-budget rejection")

# No-fractionation point-estimate counterexample. Arbitrary count units.
a6, a7, b6, b7 = 90, 10, 10, 90
z6, z7 = 80, 20
actual_source6, actual_source7 = 70, 0
assert 0 <= z6-actual_source6 <= b6
assert 0 <= z7-actual_source7 <= b7
lower = max(0,z6-b6)+max(0,z7-b7)
upper = min(z6,a6)+min(z7,a7)
assert (lower,upper) == (70,90)
point_estimate = (z6-0.1*(z6+z7))/(0.9-0.1)
assert point_estimate == 87.5 and actual_source6+actual_source7 == 70

checks = {
    "status": "PASS",
    "scope": "Exact finite conservation consistency only; no chemical response simulation or hardware test",
    "one_isotope_integer_cases": len(cases),
    "two_isotope_integer_cases": paired_cases,
    "inconsistent_budget_rejections": inconsistent,
    "no_fractionation_point_counterexample": {
        "units": "arbitrary atom-count units, mathematical counterexample only",
        "source_inputs": [a6,a7],
        "nonfeed_inputs": [b6,b7],
        "observed_product": [z6,z7],
        "actual_source_product": [actual_source6,actual_source7],
        "sharp_source_total_interval": [lower,upper],
        "invalid_unqualified_two_endmember_point": point_estimate,
    },
    "finite_routing_family": {
        "ordered_fraction_count": 15,
        "suffix_count_including_empty": 16,
        "hardware_evidence": "NOT_RUN"
    }
}
(BASE/"conservation_checks.json").write_text(json.dumps(checks,indent=2)+"\n")

manifest=json.loads((SOURCE/"SOURCE_SUBSET_MANIFEST.json").read_text())
hash_checks=[]
for record in manifest["files"]:
    p=SOURCE/record["path"]
    digest=hashlib.sha256(p.read_bytes()).hexdigest()
    assert digest == record["sha256"]
    assert p.stat().st_size == record["bytes"]
    hash_checks.append({"path":record["path"],"sha256":digest,"bytes":p.stat().st_size,"status":"match"})

folder=SOURCE/"figure 05/raw data/Li uptake experiments/Figure 5A"
p=openpyxl.load_workbook(folder/"processing.xlsx",read_only=True,data_only=False).active
cycles=[]
for k in range(1,11):
    r=2*k
    cc_initial,cc_final=p[f"B{r}"].value,p[f"B{r+1}"].value
    pu_initial,pu_final=p[f"H{r}"].value,p[f"H{r+1}"].value
    cycles.append({
        "cycle":k,
        "CC":{"initial_cell":f"B{r}","final_cell":f"B{r+1}","initial_ppm":cc_initial,"final_ppm":cc_final,
              "formula_cell":f"D{r+1}","formula":p[f"D{r+1}"].value,"formula_implied_mg_g":(cc_final-cc_initial)*0.1/0.158},
        "pulse":{"initial_cell":f"H{r}","final_cell":f"H{r+1}","initial_ppm":pu_initial,"final_ppm":pu_final,
                 "formula_cell":f"J{r+1}","formula":p[f"J{r+1}"].value,"formula_implied_mg_g":(pu_final-pu_initial)*0.1/0.158}
    })
cc_sum=sum(r["CC"]["formula_implied_mg_g"] for r in cycles)
pu_sum=sum(r["pulse"]["formula_implied_mg_g"] for r in cycles)
raw=openpyxl.load_workbook(folder/"raw data.xlsx",read_only=True,data_only=False).active
raw_headers=[[v.value for v in row] for row in list(raw.iter_rows(min_row=1,max_row=2))]

readback={
    "source_manifest":str(SOURCE/"SOURCE_SUBSET_MANIFEST.json"),
    "source_bytes":hash_checks,
    "processing_workbook":str(folder/"processing.xlsx"),
    "processing_sheet":p.title,
    "raw_workbook":str(folder/"raw data.xlsx"),
    "raw_sheet":raw.title,
    "raw_headers":raw_headers,
    "cycles":cycles,
    "summary":{
        "CC_formula_implied_sum_mg_g":cc_sum,
        "pulse_formula_implied_sum_mg_g":pu_sum,
        "pulse_CC_ratio":pu_sum/cc_sum,
        "CC_retention_formula":p["D23"].value,
        "pulse_retention_formula":p["J23"].value,
        "CC_own_initial_retention_percent":100*cycles[-1]["CC"]["formula_implied_mg_g"]/cycles[0]["CC"]["formula_implied_mg_g"],
        "pulse_own_initial_retention_percent":100*cycles[-1]["pulse"]["formula_implied_mg_g"]/cycles[0]["pulse"]["formula_implied_mg_g"],
    },
    "limitations":[
        "Original source files only read; no workbook recalculation or modification",
        "Formula conversion preserved, not validated as an absolute mass calibration",
        "Readback does not validate dilution, volume, sampling, electrode mass or source origin",
        "Elemental wavelength channels do not supply isotope trajectories",
        "Figure4 electrical and Figure5 lithium experiments were not joined",
        "Hash checks establish integrity against supplied subset manifest, not complete Zenodo acquisition",
    ],
}
(BASE/"source_readback.json").write_text(json.dumps(readback,indent=2)+"\n")
print(json.dumps({"checks":checks,"readback_summary":readback["summary"],"original_members_hash_matched":len(hash_checks)},indent=2))

