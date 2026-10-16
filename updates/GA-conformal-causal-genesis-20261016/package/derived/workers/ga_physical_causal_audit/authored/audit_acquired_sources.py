#!/usr/bin/env python3
"""Independent read-only audit of FZ native originals and integrated authored bytes.

Writes derived aggregates and receipts, never edits any source workbook. No record
is labelled raw GC chemistry, a randomized intervention, or new experiment.
"""
from datetime import datetime, timezone
from pathlib import Path
import hashlib
import json
import math
import zipfile
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
ULTRA = Path("/Users/harjas/Documents/ChatGPT/ultra")
FZ = ULTRA / "research/2026-10-16-positive-memory-genesis-v1/workers/fz_catalysis_evidence_acquisition"
EXT = FZ / "external_excluded_from_publication"
INTEGRATED = ULTRA / "work/astra-positive-memory-20261016/updates/FZ-positive-memory-genesis-20261016/package/derived/workers/fz_catalysis_evidence_acquisition"
ATTACHMENT = Path("/Users/harjas/.codex/attachments/ede43298-e149-43db-99fe-96daeff0d21c/Pasted text.txt")
REC = ROOT / "receipts"
REC.mkdir(parents=True, exist_ok=True)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def flux(x, s):
    ethane = x * (1.0 / s - 1.0)
    return {"X": x, "S_paper": s, "ethane_over_Ain": ethane,
            "net_ethylene_over_Ain_C2_only": x - ethane,
            "net_ethylene_selectivity_C2_only": 2.0 - 1.0 / s}


def header(values):
    return [[v if isinstance(v, (float, int, str)) else str(v) for v in row[:10]] for row in values[:3]]


inventory = json.loads((FZ / "receipts/public_source_inventory.json").read_text())["sources"]
receipts = []
workbook_summaries = []
for path in sorted(EXT.glob("*.xlsx")):
    inv = next(r for r in inventory if r["name"] == path.name)
    digest = sha(path)
    assert digest == inv["sha256"]
    with zipfile.ZipFile(path) as archive:
        assert "xl/workbook.xml" in archive.namelist()
        xmls = [n for n in archive.namelist() if n.startswith("xl/worksheets/sheet") and n.endswith(".xml")]
        formula_count = sum(archive.read(n).count(b"<f") for n in xmls)
    receipts.append({"file_name": path.name, "source_path": str(path), "sha256": digest,
                     "http_acquired_at_utc": inv["retrieved_at_utc"], "http_status": inv["status_code"],
                     "requested_url": inv["requested_url"], "valid_xlsx_container": True})
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    workbook_summaries.append({"file": path.name, "formula_count": formula_count,
        "sheets": [{"name": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                    "header": header(list(ws.iter_rows(max_row=3, values_only=True)))} for ws in wb]})
    wb.close()

wb = openpyxl.load_workbook(EXT / "final_Fig2_MOESM20.xlsx", read_only=True, data_only=True)
ws = wb["Figure 2 b"]
endpoints = {}
for name, row in [("no_voltage",4),("static_positive",10),("dynamic_10s_0p1s",11)]:
    endpoints[name] = {"source_cells": f"Figure 2 b!B{row}:D{row}", "source_label": ws.cell(row,2).value,
                       **flux(ws.cell(row,3).value/100.0, ws.cell(row,4).value/100.0)}
grid = {}
for sname in ("Figure 2 c", "Figure 2 d"):
    ws = wb[sname]
    grid[sname] = [{"row":r[3].row,"tp_s":r[1].value,"tm_s":r[2].value,"stored_percent":r[3].value}
                   for r in ws.iter_rows(min_row=2) if r[1].value == 10 and r[2].value in (0.1,1)]
wb.close()
preprint = {name:flux(x,s) for name,x,s in [("no_voltage",.894,.351),("static_positive",.808,.477),("dynamic_10s_0p1s",.810,.805)]}
ratios = {}
for version, points in [("preprint_v1_rounded",preprint),("final_source_cells",endpoints)]:
    ratio = points["dynamic_10s_0p1s"]["ethane_over_Ain"] / points["static_positive"]["ethane_over_Ain"]
    ratios[version] = {"ratio":ratio,"restricted_memoryless_bound":10.0/10.1,
                      "positive_hold_mean_upper_bound":1.01*ratio,
                      "suppression_equivalent_seconds_not_lifetime":10.0-10.1*ratio}

wb = openpyxl.load_workbook(EXT / "final_SIFig7_MOESM8.xlsx", read_only=True, data_only=True)
ws = wb["Figure 7 b"]
rows = np.array([[r[1],r[2],r[3]] for r in ws.iter_rows(min_row=2,values_only=True)],dtype=float)
t,v,i=rows.T
integral = np.trapezoid
current = {"source":"Figure 7 b!B2:D55000","waveform":"10 s +1.5 V / 1 s -0.1 V; not optimum 0.1 s",
           "samples":len(t),"time_first_last_s":[t[0],t[-1]],"median_dt_s":float(np.median(np.diff(t))),
           "voltage_min_max_V":[v.min(),v.max()],"current_min_max_A":[i.min(),i.max()],
           "signed_charge_C":float(integral(i,t)),"absolute_charge_C":float(integral(abs(i),t)),
           "signed_recorded_VI_J":float(integral(v*i,t)),"absolute_recorded_VI_J":float(integral(abs(v*i),t)),
           "negative_range_extreme_sample_count":int(np.sum(i==i.min())),
           "positive_range_extreme_sample_count":int(np.sum(i==i.max()))}
take=t>=3
current["tail_3s_mean_mA"]=float(integral(i[take],t[take])/(t[-1]-3)*1000)
wb.close()

wb=openpyxl.load_workbook(EXT/"final_Fig6_MOESM24.xlsx",read_only=True,data_only=True)
spectra=[]
for ws in wb:
    vals=list(ws.values)
    h=next(row for row in vals[:3] if "t / s" in row)
    columns=[j for j,x in enumerate(h) if isinstance(x,(float,int))]
    times=[h[j] for j in columns]
    spectra.append({"sheet":ws.title,"spectrum_count":len(times),"first_last_s":[min(times),max(times)],
                    "intervals_s":sorted(set(np.diff(times).tolist()))})
wb.close()

wb=openpyxl.load_workbook(EXT/"final_SIFig11_MOESM12.xlsx",read_only=True,data_only=True)
ws=wb["Figure 11 c"]
foil_control={"source":"Figure 11 c!B5:D5","condition":ws["B5"].value,
              "C_header":ws["C1"].value,"D_header":ws["D1"].value,
              "C_C2H4_to_C2H6_mol_min":ws["C5"].value,
              "D_C2H2_to_C2H4_mol_min":ws["D5"].value,
              "source_conflict":"SI physical page 22 text reverses the rounded 0.67/0.71 product assignment; plot on physical page 23 agrees with workbook."}
wb.close()

integrated_comparison=[]
for rel in ["authored/REPORT.md","authored/analysis_results.json","authored/analyze_originals.py","receipts/public_source_inventory.json"]:
    a,b=FZ/rel,INTEGRATED/rel
    integrated_comparison.append({"relative_path":rel,"worker_sha256":sha(a),"integrated_sha256":sha(b),"equal":sha(a)==sha(b)})
assert all(x["equal"] for x in integrated_comparison)

result={"scope":"read-only source audit and conditional numerical reconstruction, no new experiment",
        "audited_at_utc":datetime.now(timezone.utc).isoformat(),
        "integrated_checkout_pin_as_supplied_by_parent":"f80a8f15137575133ad2c9af975513b32e8fa93e",
        "new_report_sha256":sha(ATTACHMENT),"final_endpoints":endpoints,"preprint_v1_rounded_endpoints":preprint,
        "version_separated_ratios":ratios,"final_grid_specific_conditions":grid,
        "recorded_current":current,"final_fig6_sampling":spectra,"foil_12h_control":foil_control,
        "additional_source_files_sha256":{name:sha(EXT/name) for name in ("final_SI.pdf","preprint_v1_live.pdf","final_Fig6_caption.html")},
        "integrated_authored_byte_comparison":integrated_comparison,"workbook_inventory":workbook_summaries}
(ROOT/"authored/source_audit_results.json").write_text(json.dumps(result,indent=2)+"\n")
(REC/"FZ_original_source_hash_checks.json").write_text(json.dumps(receipts,indent=2)+"\n")
print(json.dumps({"ratio_final":ratios["final_source_cells"]["ratio"],"ratio_preprint":ratios["preprint_v1_rounded"]["ratio"],
                  "acquired_workbooks":len(receipts),"current_samples":current["samples"],"integrated_authored_equal":True},indent=2))
