"""Read-only numerical acquisition from the AgRu primary source data.

No fitted catalyst property is invented. Interpolated LSV values are labelled
as such; half-cell data are never substituted for measured full-cell voltages.
Thermodynamic calculations use common, restored carrier/product states.
Run with the bundled Python (openpyxl is used only to read source XLSX).
"""
from pathlib import Path
from math import log
import hashlib
import json
import openpyxl

HERE = Path(__file__).resolve().parent
SOURCE = HERE / 'sources/41467_2026_71299_MOESM5_ESM.xlsx'
F = 96485.33212                 # C mol-1
M_NH3 = 0.017031                # kg mol-1, same as paper TEA
M_H2 = 0.002016                 # kg mol-1
M_N = 14.007                    # g mol-1
G_H2 = 237.141                  # kJ mol-1 H2O(l) -> H2 + .5 O2
G_NH3 = -16.367                 # kJ mol-1, JANAF NH3(g), 298.15K,1bar
G_TARGET = G_NH3 + 1.5 * G_H2

def kwhkg(kjmol):
    return kjmol / (3600 * M_NH3)

def interp_first(points, target):
    """Linear interpolation in the first adjacent bracket, not extrapolation."""
    for (x0,y0),(x1,y1) in zip(points, points[1:]):
        if min(x0,x1) <= target <= max(x0,x1) and x0 != x1:
            return y0 + (y1-y0)*(target-x0)/(x1-x0)
    return None

def pairs(sheet, xc, yc, first=3):
    out=[]
    for row in sheet.iter_rows(min_row=first, values_only=True):
        x,y=row[xc-1],row[yc-1]
        if isinstance(x,(int,float)) and isinstance(y,(int,float)):
            out.append((x,y))
    return out

def regress_origin(points):
    return sum(x*y for x,y in points)/sum(x*x for x,y in points)

w=openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
out={
    'source_url':'https://www.nature.com/articles/s41467-026-71299-x',
    'source_xlsx_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'nature_of_results':'source-reported data plus labelled arithmetic; no replication or new experiment',
    'thermodynamic_sources':{
        'water_liquid':'https://janaf.nist.gov/tables/H-063.html',
        'ammonia_gas':'https://janaf.nist.gov/tables/H-083.html',
        'common_standard_state':'298.15K,1bar ideal gases,pure liquid water; no air separation or NH3 liquefaction credit'
    },
    'constants':dict(F_C_mol=F,M_NH3_kg_mol=M_NH3,M_H2_kg_mol=M_H2,
                     G_water_splitting_kJ_mol=G_H2,G_ammonia_air_water_kJ_mol=G_TARGET),
}

# Actual full-cell current/voltage curves: consumed voltage is negative in
# galvanic operation. Unlike a half-cell RHE value, this includes the anode.
s=w['Figure  6c']
curves={}
for c,xc,yc in [(10,1,2),(50,4,5),(100,7,8),(500,10,11),(1000,13,14)]:
    p=pairs(s,xc,yc)
    power=[(j*(-v),j,v) for j,v in p if j>=0 and v<=0]
    peak=max(power)
    curves[str(c)]={
        'concentration_mM':c,'data_region':f'Figure  6c columns {xc},{yc},rows3:{s.max_row}',
        'LSV_not_steady_state':True,
        'interpolated_consumed_voltage_V_at_j_mA_cm2':{
            str(j):interp_first(p,j) for j in (10,50,100,200,500,1000)},
        'peak_export_power_in_LSV_mW_cm2':peak[0],
        'at_j_mA_cm2':peak[1],'consumed_voltage_at_peak_V':peak[2],
        'j_at_zero_consumed_voltage_interpolated_mA_cm2':interp_first([(v,j) for j,v in p],0)
    }
out['paired_nitrate_HOR_source_curves']=curves

# The time series contains successive renewed electrolyte batches; sum only
# positive dt within each series, never connect across renewal discontinuities.
s=w['Figure  6d']
weighted_v=duration=0.0
segments=[]
for xc in (4,6,8,10,12,14,16,18,20):
    p=pairs(s,xc,xc+1,first=2)
    dt_sum=integ=0.0
    for (t0,v0),(t1,v1) in zip(p,p[1:]):
        dt=t1-t0
        if dt>0:
            dt_sum+=dt;integ+=dt*(v0+v1)/2
    if dt_sum:
        segments.append({'time_column':xc,'hours':dt_sum,'mean_consumed_voltage_V':integ/dt_sum})
        duration+=dt_sum;weighted_v+=integ
fe=pairs(s,1,2,first=2)
out['paired_nitrate_100h']={
    'integrated_observed_duration_h':duration,
    'time_weighted_consumed_voltage_V':weighted_v/duration,
    'segments':segments,
    'FE_points_percent':[y for x,y in fe],
    'FE_point_arithmetic_mean_percent':sum(y for x,y in fe)/len(fe),
    'source_scope':'0.25cm2 cathode;100mM nitrate,1M KOH;200mA/cm2;9 electrolyte renewals;not closed carrier recycling',
}

# Nitrite-specific measured half-cell LSV, with/without HER-background removal.
s=w['Supplementary Figure 20']
ni_raw=pairs(s,19,20,first=4)  # Fig.S20c, AgRu-15, S:T
ni_sub=pairs(s,35,36,first=4)  # Fig.S20e, AgRu-15, AI:AJ
out['nitrite_half_cell']={
    'conditions':'10mM KNO2+1M KOH;0.6cm2;source-reported iR correction;not a full-cell power measurement',
    'interpolated_j_mA_cm2_at_E_RHE_V':{
        str(e):{'raw':interp_first(ni_raw,e),'HER_background_subtracted':interp_first(ni_sub,e)}
        for e in (0.3,0.2,0.1,0.05,0.0,-0.05)
    },
    'onset_E_RHE_at_2mA_cm2':w['Figure  3c']['F4'].value,
    'Tafel_slope_mV_dec':w['Figure  3c']['G4'].value,
    'Tafel_prediction_E_RHE_at_200mA_cm2_V':w['Figure  3c']['F4'].value-w['Figure  3c']['G4'].value/1000*log(200/2,10),
    'Tafel_scope':'prediction only; may leave fitted kinetic region, not an acquired operating value',
}

# Direct nitrite concentration/product time series (source S21d). This gives
# a measured-concentration-derived equivalent current, not a measured FE.
s=w['Supplementary Figure 21']
initial=s['S4'].value;remain=s['S9'].value;prod=s['T9'].value
mol_nh3=prod*30e-6/M_N  # ug/mL *30mL /1e6 ug/g /(g/mol)
avg_j=mol_nh3*6*F/(100*60)/0.6*1000
sum_residual=max(abs(s[f'S{r}'].value+s[f'T{r}'].value-initial) for r in range(4,10))
out['nitrite_concentration_time_series']={
    'range':'Supplementary Figure 21,R4:T9;0.1V_RHE;30mL;0.6cm2',
    'N_initial_ug_mL':initial,'nitrite_N_at_100min_ug_mL':remain,
    'NH3_N_at_100min_ug_mL':prod,'converted_fraction':1-remain/initial,
    'NH3_equivalent_average_partial_j_mA_cm2':avg_j,
    'max_abs_closure_residual_ug_mL':sum_residual,
    'caution':'S+T sum is exact to printed precision at all six times. It cannot count as independent multiproduct nitrogen-balance validation; no inference excluding hydroxylamine/N2 is made.'
}

# Exact material supply lower bound, j=nF k_m C, independent of an invented D.
out['nitrite_supply_at_200mA_cm2']={
    str(c):{'concentration_mM':c,'minimum_km_m_s':2000/(6*F*c)}
    for c in (1,10,100,1000)
}

# Optimistic, fully reversible non-electrolyser redox circuit. The loop must
# still pay actual electrolyser dissipation once for EVERY hydrogen molecule.
# n=(z+3)/2; W_total >= G_target+n*(w_H2-G_H2), if all other operations are
# reversible and all common stream states are restored. Extra heat/exergy
# sources, non-unit yield, pressure and separation require explicit terms.
rows=[]
for wh in (40,45,50,55):
    kj=wh*M_H2*3600
    for z,name in ((0,'direct_N2'),(2,'NO'),(3,'nitrite'),(5,'nitrate')):
        n=(z+3)/2
        rows.append({'w_H2_kWh_kgH2_sensitivity_input':wh,'N_oxidation_state':z,'carrier':name,
                     'H2_mol_per_NH3':n,
                     'ideal_other_stages_floor_kWh_kgNH3':kwhkg(G_TARGET+n*(kj-G_H2)),
                     'extra_floor_vs_direct_kWh_kgNH3':kwhkg(z/2*(kj-G_H2))})
out['conditional_ideal_loop_floor_sensitivity']=rows

# Acquired paper benchmark, source SI57-59. 50kWh/kgH2 is explicitly a
# sensitivity input, not a property measured in that paper or a plant quote.
# No atmospheric fixation term has yet been added. Captured electrolyte base
# and product states must be restored before a total-process comparison.
point=[]
for wh in (40,45,50,55):
    for name,z,fe,eout in [('acquired_nitrate_point',5,.81,.1),
                          ('hypothetical_nitrite_same_voltage_unit_FE',3,1,.1)]:
        ne=z+3;h2mol=ne/(2*fe);h2kgkg=h2mol*M_H2/M_NH3
        recovered=kwhkg(ne*F*eout/1000/fe)
        point.append({'name':name,'source_or_assumption':'source nitrate point; nitrite point is optimistic scenario, NOT measured full cell' if z==3 else 'SI57-59 measured/assumed TEA point',
                      'w_H2_kWh_kgH2_sensitivity_input':wh,
                      'H2_kg_kgNH3':h2kgkg,'H2_electricity_kWh_kgNH3':h2kgkg*wh,
                      'electrical_return_kWh_kgNH3':recovered,
                      'partial_energy_excluding_fixation_and_separation_kWh_kgNH3':h2kgkg*wh-recovered,
                      'paper_stripping_proxy_kWh_kgNH3':5.6*14/M_NH3/1000})
out['source_point_and_nitrite_scenario']=point

# Next justified construction: feed NO directly and replace HOR with OER,
# avoiding the extra H2 electrolyser. Same gas-NH3/liquid-water standard states
# are used for the reversible limit, never a mismatched NH4+/pH potential.
G_NO=86.600  # JANAF N-005; older evaluated table used consistently here.
G_NO_HOR=G_NH3-G_H2-G_NO
G_NO_OER=G_TARGET-G_NO
out['next_NO_construction']={
    'NO_standard_G_source':'https://janaf.nist.gov/tables/N-005.html',
    'HOR_net':'NO(g)+2.5H2(g)->NH3(g)+H2O(l)',
    'HOR_reversible_release_kJ_mol':-G_NO_HOR,
    'HOR_reversible_cell_V_same_product_states':-G_NO_HOR*1000/(5*F),
    'OER_net':'NO(g)+1.5H2O(l)->NH3(g)+1.25O2(g)',
    'OER_reversible_input_kJ_mol':G_NO_OER,
    'OER_reversible_cell_V_same_product_states':G_NO_OER*1000/(5*F),
    'source_OER_point':'Bunea2023,1.9V input,93%NH3FE,NO/He feed;product collected in acid,so full gasNH3 recovery is NOT acquired',
    'source_OER_point_electricity_kWh_kgNH3':kwhkg(5*F*1.9/1000/.93),
    'oxidation_power_boundary_cases':[
        {'oxidation_MJ_molN':p,'label':label,
         'oxidation_only_kWh_kgNH3':kwhkg(p*1000),
         'NO_OER_ideal_other_stages_kWh_kgNH3':kwhkg(p*1000+G_NO_OER),
         'NO_OER_source_point_plus_oxidation_kWh_kgNH3':kwhkg(p*1000+5*F*1.9/1000/.93),
         'NO_HOR_ideal_return_50kWh_H2_partial_kWh_kgNH3':kwhkg(p*1000+2.5*50*M_H2*3600+G_NO_HOR)}
        for p,label in ((.42,'reported absorbed plasma only;generous non-wall boundary'),
                        (.42*.76/.1,'same reported plasma yield normalized to reported applied power;approximate rounded powers'))
    ],
    'necessary_source_NO_OER_voltage_for_12p5kWh_total_with_absorbed_plasma_V':
        (12.5*M_NH3*3600-420)*.93*1000/(5*F),
    'scope':'NO/HOR perfect-return limits are thermodynamic ceilings, NOT demonstrated operating values. Product/separation,plasma gas dilution,air conditioning,ion cycling and auxiliaries still require payment.'
}

out['matching_HB_and_loaded_voltage_budgets']={
    'definition':'W_HB(kJ/molNH3)=1.5*w_H2+H_HB,where H_HB is actual net loop/ASU/separation/control/storage overhead at same scale and heat quality',
    '50kWhH2_common_HB_H2_only_kWh_kgNH3':1.5*M_H2/M_NH3*50,
    'nitrate_paper_point_min_HB_overhead_to_be_better_before_fixation_or_product_separation_kWh_kgNH3':
        (4/.81-1.5)*M_H2/M_NH3*50-kwhkg(8*F*.1/1000/.81),
    'nitrite_optimistic_0p1V_unitFE_min_HB_overhead_to_be_better_before_fixation_or_separation_kWh_kgNH3':
        1.5*M_H2/M_NH3*50-kwhkg(6*F*.1/1000),
    'required_nitrite_exported_voltage_50kWhH2_no_fixation_or_separation':{
        str(b):(3*50*M_H2*3600-b*M_NH3*3600)*1000/(6*F) for b in (10.5,12.5,15)
    },
    'scope':'HB energy budgets 10.5/12.5/15 are sensitivity targets,not acquired named-plant quotes. Source isolated33.9GJ/t is not combined with50kWh/kgH2 without identifying its configuration.'
}

(HERE/'operation_calculations.json').write_text(json.dumps(out,indent=2)+'\n')
print(json.dumps({k:v for k,v in out.items() if k not in ('paired_nitrate_HOR_source_curves','paired_nitrate_100h')},indent=2))
print('FULL CELL LSV SUMMARY')
for key,v in curves.items():
    print(key,'mM: V_in(j=200)=',v['interpolated_consumed_voltage_V_at_j_mA_cm2']['200'],
          'max_export_mW/cm2=',v['peak_export_power_in_LSV_mW_cm2'],
          'j_zeroV=',v['j_at_zero_consumed_voltage_interpolated_mA_cm2'])
print('100h time-weighted consumed voltage',weighted_v/duration,'over',duration,'h')
